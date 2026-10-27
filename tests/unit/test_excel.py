# SPDX-License-Identifier: Apache-2.0
"""Outil fourni ``loom.excel`` : description sémantique, style maison, validation, recalcul.

L'outil est construit par sa fabrique sur des dossiers temporaires ; ses appels
passent par ``invoke`` et ``assess_call``, les deux chemins du moteur. Les
classeurs écrits sont relus avec openpyxl.
"""

import subprocess
import sys
from pathlib import Path
from typing import Any, Final, cast

import pytest

pytest.importorskip("openpyxl", reason="extra 'office' absent")

from conftest import MODEL, ConfigFactory, demo_agent
from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font
from pydantic import JsonValue, ValidationError

import loom_ia.tools_builtin.excel.recalc as recalc_module
from loom_ia.access.cli import main
from loom_ia.adapters.stores import InMemoryEventStore
from loom_ia.config import ConfigError, load_config
from loom_ia.core.events import ToolCompleted
from loom_ia.core.model import DEFAULT_TENANT, RunId, RunStatus, SessionId
from loom_ia.core.ports import ToolContext, ToolError
from loom_ia.engine import begin_run, drive
from loom_ia.runtime import build_agent
from loom_ia.tools_builtin import BuiltinContext
from loom_ia.tools_builtin.excel import ExcelTool, build
from loom_ia.tools_builtin.excel.arguments import Problems, SheetSpec
from loom_ia.tools_builtin.excel.formulas import check_formula, quoted, sheet_name_problem
from loom_ia.tools_builtin.excel.sources import read_rows
from loom_ia.tools_builtin.excel.workbook import FIXED_DATE
from loom_ia.tools_builtin.excel_params import ExcelParams

CONTEXT: Final = ToolContext(
    tenant_id=DEFAULT_TENANT,
    session_id=SessionId("s"),
    run_id=RunId("r"),
    call_id="c1",
    agent="demo",
)
JSONL: Final = {"events": {"backend": "jsonl", "path": "journal"}}


def make(tmp_path: Path, *folders: dict[str, Any], **params: Any) -> ExcelTool:
    model = ExcelParams.model_validate({"folders": list(folders), **params})
    return build(model, BuiltinContext(base_dir=tmp_path))


@pytest.fixture
def disk(tmp_path: Path) -> Path:
    for name in ("donnees", "sortie", "valide"):
        (tmp_path / name).mkdir()
    return tmp_path


@pytest.fixture
def tool(disk: Path) -> ExcelTool:
    return make(
        disk,
        {"name": "donnees", "path": "donnees", "rights": ["read"]},
        {"name": "sortie", "path": "sortie", "rights": ["write", "overwrite"]},
        {"name": "valide", "path": "valide", "rights": ["write"], "approval": {"write": "always"}},
    )


async def call(tool: ExcelTool, **arguments: Any) -> dict[str, Any]:
    output = await tool.invoke(dict(arguments), CONTEXT)
    assert isinstance(output.data, dict)
    return output.data


async def refused(tool: ExcelTool, match: str, **arguments: Any) -> None:
    with pytest.raises(ToolError, match=match):
        await tool.invoke(dict(arguments), CONTEXT)


def rgb(found: Cell) -> str | None:
    """La couleur du texte de la cellule, ``None`` sans couleur."""
    color = cast(Font, found.font).color
    return None if color is None else str(color.rgb)


def aligned(found: Cell) -> Alignment:
    return cast(Alignment, found.alignment)


def cell(path: Path, sheet: str, address: str) -> Cell:
    found = load_workbook(path)[sheet][address]
    assert isinstance(found, Cell)
    return found


SUIVI: Final[dict[str, Any]] = {
    "name": "Suivi",
    "columns": [
        {"header": "Client", "type": "texte"},
        {"header": "Date devis", "type": "date"},
        {"header": "Montant HT", "type": "montant_eur"},
        {"header": "Relances", "type": "entier"},
    ],
    "rows": [
        {"cells": ["Dupont", "2026-10-02", 1250.5, 2]},
        {"cells": ["Martin", "2026-10-09", 830, 0]},
    ],
    "totals": {"Montant HT": "somme", "Relances": "somme"},
}
HYPOTHESES: Final[dict[str, Any]] = {
    "name": "Hypothèses",
    "columns": [{"header": "Paramètre"}, {"header": "Valeur", "type": "pourcentage"}],
    "rows": [
        {
            "cells": ["Taux de conversion visé", 0.15],
            "role": "hypothese",
            "note": "Source : objectif fixé par l'utilisateur",
        }
    ],
}


# --- Ce que voit le modèle ------------------------------------------------------------


def test_the_spec_names_the_folders_and_the_vocabulary(tool: ExcelTool) -> None:
    spec = tool.spec
    assert spec.name == "excel"
    assert "- donnees : read\n" in spec.description
    assert "- valide : write (attend une approbation humaine : write)" in spec.description
    assert "montant_eur" in spec.description and "hypothese" in spec.description
    assert (spec.approval, spec.side_effects) == ("always", "irreversible")
    assert "sheets" in str(spec.input_schema)


def test_settings_lines_for_validate(tool: ExcelTool, disk: Path) -> None:
    assert tool.settings_lines() == (
        f"donnees → {disk / 'donnees'}",
        "  read : accordé, sans approbation",
        f"sortie → {disk / 'sortie'}",
        "  write : accordé, sans approbation",
        "  overwrite : accordé, avec approbation",
        f"valide → {disk / 'valide'}",
        "  write : accordé, avec approbation",
        "fuseau : Europe/Paris",
        "recalcul : non",
        "lecture des classeurs : oui",
    )


def test_the_schema_has_no_titles(tool: ExcelTool) -> None:
    """Le schéma part au modèle à chaque appel : sans les titres que Pydantic ajoute."""
    assert '"title"' not in str(tool.spec.input_schema).replace("'", '"')
    assert "Europe/Paris" in tool.spec.description
    # L'unité qu'affiche déjà le format ne se répète pas dans l'en-tête (essai réel du 26/10).
    assert "ne pas la répéter dans l'en-tête" in tool.spec.description
    assert "pas « Montant HT (€) »" in str(tool.spec.input_schema)


# --- Paramètres -------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("params", "message"),
    [
        (
            {"folders": [{"name": "s", "path": "s", "rights": ["write", "delete"]}]},
            "le droit 'delete' ne sert pas à loom.excel",
        ),
        # delete dans approval : refusé pour excel, et non « l'ajouter à rights ».
        (
            {
                "folders": [
                    {
                        "name": "s",
                        "path": "s",
                        "rights": ["write"],
                        "approval": {"delete": "always"},
                    }
                ]
            },
            "dossier s : le droit 'delete' ne sert pas à loom.excel",
        ),
        (
            {"folders": [{"name": "s", "path": "s", "rights": {"write": "never"}}]},
            "dossier 's' : 'rights' en table droit → approbation n'est plus accepté",
        ),
        ({"folders": []}, "au moins un dossier est attendu"),
        (
            {
                "folders": [
                    {"name": "s", "path": "s", "rights": ["write"]},
                    {"name": "s", "path": "t", "rights": ["read"]},
                ]
            },
            "nom de dossier donné deux fois : s",
        ),
        (
            {"folders": [{"name": "s", "path": "s", "rights": ["write"]}], "soffice": "/x"},
            "'soffice' ne sert qu'avec recalc: true",
        ),
        (
            {"folders": [{"name": "s", "path": "s", "rights": ["write"]}], "recalc_timeout": 5},
            "'recalc_timeout' ne sert qu'avec recalc: true",
        ),
        (
            {"folders": [{"name": "s", "path": "s", "rights": ["write"]}], "timezone": "Mars/Base"},
            "fuseau inconnu : 'Mars/Base'",
        ),
    ],
)
def test_params_refused(params: dict[str, Any], message: str) -> None:
    with pytest.raises(ValidationError, match=message):
        ExcelParams.model_validate(params)


def test_recalc_needs_libreoffice(disk: Path) -> None:
    with pytest.raises(ValueError, match="LibreOffice est introuvable"):
        make(
            disk,
            {"name": "sortie", "path": "sortie", "rights": ["write"]},
            recalc=True,
            soffice="/nulle/part/soffice",
        )


# --- Écrire -----------------------------------------------------------------------------


async def test_writes_a_formatted_workbook(tool: ExcelTool, disk: Path) -> None:
    done = await call(tool, path="sortie/relances.xlsx", sheets=[HYPOTHESES, SUIVI])
    target = disk / "sortie" / "relances.xlsx"
    assert done["path"] == "sortie/relances.xlsx"
    assert (done["written"], done["replaced"]) == (True, False)
    assert done["sheets"] == [
        {"name": "Hypothèses", "rows": 1, "formulas": 0},
        {"name": "Suivi", "rows": 2, "formulas": 2},
    ]
    workbook = load_workbook(target)
    assert workbook.sheetnames == ["Hypothèses", "Suivi"]
    suivi = workbook["Suivi"]
    assert suivi.freeze_panes == "A2"
    assert suivi.auto_filter.ref == "A1:D3"
    header = suivi["A1"]
    assert header.font is not None and header.font.b
    assert suivi["B2"].number_format == "mm-dd-yy"
    assert suivi["C2"].number_format == '#,##0.00 "€"'
    assert (suivi["A4"].value, suivi["C4"].value, suivi["D4"].value) == (
        "Total",
        "=SUM(C2:C3)",
        "=SUM(D2:D3)",
    )
    assert suivi["C4"].font is not None and suivi["C4"].font.b
    hypotheses = workbook["Hypothèses"]
    value = hypotheses["B2"]
    assert value.value == 0.15 and value.number_format == "0.0%"
    assert rgb(value) == "000000FF"
    assert value.comment is not None and "Source" in value.comment.text
    # Le libellé d'une ligne d'hypothèse reste noir.
    label = hypotheses["A2"]
    assert label.font is not None and (
        label.font.color is None or label.font.color.rgb != "000000FF"
    )
    assert hypotheses["A4"].value == "Bleu sur fond jaune : hypothèse (sa source en commentaire)."
    width = suivi.column_dimensions["A"].width
    assert width is not None and 8 <= width <= 60


PRELEVEMENTS: Final[dict[str, Any]] = {
    "name": "Prélèvements",
    "columns": [
        {"header": "Date", "type": "date"},
        {"header": "Libellé"},
        {"header": "Montant", "type": "montant_eur"},
        {"header": "Doublon signalé"},
    ],
    "rows": [
        {
            "cells": [
                "2026-10-21",
                "Commande 5521 du 20/10",
                89.9,
                "Oui : même montant que PRL-91307",
            ]
        },
        {
            "cells": [
                "2026-10-21",
                "Commande 5521 du 20/10",
                89.9,
                "Oui : même montant que PRL-91306",
            ]
        },
    ],
    "totals": {"Montant": "somme"},
}


async def test_the_total_is_labelled_in_the_first_text_column(tool: ExcelTool, disk: Path) -> None:
    """Une date en première colonne ne laisse plus le total sans libellé (essai réel du 26/10)."""
    await call(tool, path="sortie/p.xlsx", sheets=[PRELEVEMENTS])
    target = disk / "sortie/p.xlsx"
    assert cell(target, "Prélèvements", "A4").value is None
    assert cell(target, "Prélèvements", "B4").value == "Total"
    assert cell(target, "Prélèvements", "C4").value == "=SUM(C2:C3)"
    everything = {**PRELEVEMENTS, "columns": [{"header": "Montant", "type": "montant_eur"}]}
    everything["rows"] = [{"cells": [10]}, {"cells": [20]}]
    await call(tool, path="sortie/q.xlsx", sheets=[everything])
    # Toutes les colonnes ont leur agrégat : pas de place pour un libellé.
    assert cell(disk / "sortie/q.xlsx", "Prélèvements", "A4").value == "=SUM(A2:A3)"


async def test_alignment_keeps_columns_apart(tool: ExcelTool, disk: Path) -> None:
    """Dates centrées, texte en retrait : rien ne colle au nombre aligné à droite d'à côté."""
    await call(tool, path="sortie/p.xlsx", sheets=[PRELEVEMENTS, HYPOTHESES])
    target = disk / "sortie/p.xlsx"
    date = aligned(cell(target, "Prélèvements", "A2"))
    assert date.horizontal == "center"
    text = aligned(cell(target, "Prélèvements", "D2"))
    assert (text.horizontal, text.indent) == ("left", 1)
    assert aligned(cell(target, "Prélèvements", "C2")).horizontal is None
    legend = aligned(cell(target, "Hypothèses", "A4"))
    assert legend.indent == 0


async def test_each_sheet_prints_one_page_wide(tool: ExcelTool, disk: Path) -> None:
    """Sans réglage, une colonne de trop partait seule sur une deuxième page (essai réel)."""
    wide = {
        **PRELEVEMENTS,
        "name": "Large",
        "columns": [*PRELEVEMENTS["columns"], {"header": "Commentaire " + "x" * 60}],
    }
    await call(tool, path="sortie/p.xlsx", sheets=[PRELEVEMENTS, wide])
    workbook = load_workbook(disk / "sortie/p.xlsx")
    for name, orientation in (("Prélèvements", "portrait"), ("Large", "landscape")):
        sheet = workbook[name]
        setup = sheet.page_setup
        assert (setup.orientation, setup.fitToWidth, setup.fitToHeight) == (orientation, 1, 0)
        properties = sheet.sheet_properties.pageSetUpPr
        assert properties is not None and properties.fitToPage


async def test_same_description_same_file(tool: ExcelTool, disk: Path) -> None:
    """Dates fixées : un second appel identique ne réécrit rien (reprise idempotente)."""
    first = await call(tool, path="sortie/a.xlsx", sheets=[SUIVI])
    second = await call(tool, path="sortie/a.xlsx", sheets=[SUIVI])
    assert first["sha256"] == second["sha256"]
    assert (first["written"], second["written"]) == (True, False)
    other = await call(tool, path="sortie/b.xlsx", sheets=[SUIVI])
    assert other["sha256"] == first["sha256"]
    assert (disk / "sortie/a.xlsx").read_bytes() == (disk / "sortie/b.xlsx").read_bytes()
    # Deux appels dans la même seconde se ressembleraient même sans dates fixées : on les lit.
    properties = load_workbook(disk / "sortie/a.xlsx").properties
    assert (properties.created, properties.modified) == (FIXED_DATE,) * 2


async def test_replacing_needs_overwrite(tool: ExcelTool, disk: Path) -> None:
    await call(tool, path="sortie/a.xlsx", sheets=[SUIVI])
    await refused(tool, "existe déjà : overwrite=true", path="sortie/a.xlsx", sheets=[HYPOTHESES])
    done = await call(tool, path="sortie/a.xlsx", sheets=[HYPOTHESES], overwrite=True)
    assert (done["written"], done["replaced"]) == (True, True)
    assert load_workbook(disk / "sortie/a.xlsx").sheetnames == ["Hypothèses"]
    await call(tool, path="valide/a.xlsx", sheets=[SUIVI])
    await refused(
        tool, "Droit overwrite refusé", path="valide/a.xlsx", sheets=[SUIVI], overwrite=True
    )


async def test_paths_follow_local_files(tool: ExcelTool) -> None:
    await refused(tool, "s'écrit en .xlsx", path="sortie/a.csv", sheets=[SUIVI])
    await refused(tool, "« .. » n'est pas permis", path="sortie/../a.xlsx", sheets=[SUIVI])
    await refused(tool, "Dossier inconnu", path="ailleurs/a.xlsx", sheets=[SUIVI])
    await refused(tool, "Droit write refusé", path="donnees/a.xlsx", sheets=[SUIVI])
    source = {**SUIVI, "rows": None, "data": {"path": "sortie/x.csv"}}
    await refused(tool, "Droit read refusé", path="sortie/a.xlsx", sheets=[source])


async def test_a_link_out_of_the_folder_is_refused(tool: ExcelTool, disk: Path) -> None:
    (disk / "ailleurs").mkdir()
    (disk / "sortie" / "lien").symlink_to(disk / "ailleurs")
    await refused(tool, "hors des dossiers permis", path="sortie/lien/a.xlsx", sheets=[SUIVI])
    assert not (disk / "ailleurs" / "a.xlsx").exists()


# --- Approbation par appel -------------------------------------------------------------


def test_assess_asks_only_where_the_folder_asks(tool: ExcelTool, disk: Path) -> None:
    free = tool.assess_call({"path": "sortie/a.xlsx", "sheets": [SUIVI]})
    assert (free.approval, free.side_effects, free.idempotent) == ("never", "reversible", True)
    replace = tool.assess_call({"path": "sortie/a.xlsx", "sheets": [SUIVI], "overwrite": True})
    # En liste, overwrite prend son approbation par défaut : always.
    assert (replace.approval, replace.side_effects) == ("always", "irreversible")
    assert replace.reason == (
        "écrire le classeur sortie/a.xlsx : approbation demandée pour overwrite sur sortie/a.xlsx"
    )
    asked = tool.assess_call({"path": "valide/a.xlsx", "sheets": [SUIVI]})
    assert asked.approval == "always"
    assert asked.reason == (
        "écrire le classeur valide/a.xlsx : approbation demandée pour write sur valide/a.xlsx"
    )
    guarded = make(
        disk,
        {"name": "donnees", "path": "donnees", "rights": ["read"], "approval": {"read": "always"}},
        {"name": "sortie", "path": "sortie", "rights": ["write"]},
    )
    source = {**SUIVI, "rows": None, "data": {"path": "donnees/devis.csv"}}
    arguments: dict[str, Any] = {"path": "sortie/a.xlsx", "sheets": [source]}
    read = guarded.assess_call(arguments)
    assert (read.approval, read.reason) == (
        "always",
        "écrire le classeur sortie/a.xlsx : approbation demandée pour read sur donnees/devis.csv",
    )
    wrong = tool.assess_call({"path": "/etc/a.xlsx", "sheets": [SUIVI]})
    assert (wrong.approval, wrong.side_effects) == ("never", "none")


# --- Sources de données ----------------------------------------------------------------


async def test_reads_a_french_csv(tool: ExcelTool, disk: Path) -> None:
    (disk / "donnees" / "devis.csv").write_text(
        "Client;Date devis;Montant HT;Relances;Commentaire\n"
        "Dupont;02/10/2026;1 250,50;2;rappeler\n"
        "Martin;2026-10-09;830;0;\n",
        encoding="utf-8-sig",
    )
    sheet = {**SUIVI, "rows": None, "data": {"path": "donnees/devis.csv"}}
    done = await call(tool, path="sortie/devis.xlsx", sheets=[sheet])
    assert done["sheets"] == [{"name": "Suivi", "rows": 2, "formulas": 2}]
    target = disk / "sortie" / "devis.xlsx"
    assert cell(target, "Suivi", "C2").value == 1250.5
    assert cell(target, "Suivi", "D2").value == 2
    assert str(cell(target, "Suivi", "B2").value).startswith("2026-10-02")


async def test_reads_json_and_says_missing_columns(tool: ExcelTool, disk: Path) -> None:
    (disk / "donnees" / "devis.json").write_text(
        '[{"Client": "Dupont", "Date devis": "2026-10-02", "Montant HT": 10, "Relances": 1}]',
        encoding="utf-8",
    )
    sheet = {**SUIVI, "rows": None, "data": {"path": "donnees/devis.json"}}
    done = await call(tool, path="sortie/j.xlsx", sheets=[sheet])
    assert done["sheets"][0]["rows"] == 1
    (disk / "donnees" / "court.json").write_text('[{"Client": "Dupont"}]', encoding="utf-8")
    short = {**sheet, "data": {"path": "donnees/court.json"}}
    await refused(
        tool,
        "colonne\\(s\\) absente\\(s\\) du fichier : Date devis",
        path="sortie/k.xlsx",
        sheets=[short],
    )
    await refused(
        tool,
        "fichier introuvable",
        path="sortie/k.xlsx",
        sheets=[{**sheet, "data": {"path": "donnees/rien.json"}}],
    )


@pytest.mark.parametrize(
    ("data", "suffix", "said"),
    [
        (b"x", ".xls", "seuls les fichiers .csv, .json et .xlsx sont lus"),
        (b"\xff\xfe", ".csv", "n'est pas un fichier texte UTF-8"),
        (b"Client\n" + b'"' + b"x" * 200_000 + b'"\n', ".csv", "CSV illisible"),
        (b"{", ".json", "JSON illisible"),
        (b'{"Client": "x"}', ".json", "une liste d'objets est attendue"),
        (b'["x"]', ".json", "une liste d'objets est attendue"),
        (b'[{"Client": {"nom": "x"}}]', ".json", "valeur non scalaire pour 'Client'"),
    ],
)
def test_unreadable_sources(data: bytes, suffix: str, said: str) -> None:
    sheet = SheetSpec.model_validate({"name": "S", "columns": [{"header": "Client"}]})
    with pytest.raises(ToolError, match=said):
        read_rows(data, suffix, sheet, "donnees/f")


def test_an_empty_source_gives_no_rows() -> None:
    sheet = SheetSpec.model_validate({"name": "S", "columns": [{"header": "Client"}]})
    assert read_rows(b"", ".csv", sheet, "donnees/f") == []
    assert read_rows(b"[]", ".json", sheet, "donnees/f") == []


async def test_inline_rows_are_bounded(tool: ExcelTool) -> None:
    many = {**SUIVI, "rows": [{"cells": ["x"]}] * 201}
    await refused(
        tool, "au-delà de 200 — les écrire dans un fichier", path="sortie/a.xlsx", sheets=[many]
    )
    both = {**SUIVI, "data": {"path": "donnees/a.csv"}}
    await refused(tool, "'rows' ou 'data', pas les deux", path="sortie/a.xlsx", sheets=[both])


# --- Valeurs ---------------------------------------------------------------------------


def _one(column_type: str, *values: JsonValue) -> JsonValue:
    return {
        "name": "F",
        "columns": [{"header": "V", "type": column_type}],
        "rows": [{"cells": [value]} for value in values],
    }


async def test_values_follow_their_type(tool: ExcelTool, disk: Path) -> None:
    sheets: list[JsonValue] = [
        {
            "name": "Types",
            "columns": [
                {"header": "Année", "type": "annee"},
                {"header": "Actif", "type": "booleen"},
                {"header": "Quand", "type": "date_heure"},
                {"header": "Part", "type": "pourcentage"},
            ],
            "rows": [{"cells": [2026, True, "2026-10-25T17:50:00+00:00", "15%"]}],
        }
    ]
    done = await call(tool, path="sortie/t.xlsx", sheets=sheets)
    assert done["warnings"] == []
    target = disk / "sortie" / "t.xlsx"
    assert cell(target, "Types", "A2").value == "2026"
    assert cell(target, "Types", "B2").value == "Oui"
    # Excel ne connaît pas les fuseaux : l'heure de Paris, sans fuseau.
    assert str(cell(target, "Types", "C2").value) == "2026-10-25 18:50:00"
    assert cell(target, "Types", "D2").value == 0.15


async def test_the_timezone_is_a_parameter(disk: Path) -> None:
    tool = make(
        disk,
        {"name": "sortie", "path": "sortie", "rights": ["write"]},
        timezone="America/New_York",
    )
    sheet: dict[str, Any] = {
        "name": "T",
        "columns": [{"header": "Quand", "type": "date_heure"}],
        "rows": [{"cells": ["2026-10-25T17:50:00+00:00"]}, {"cells": ["2026-10-25 09:00"]}],
    }
    await call(tool, path="sortie/t.xlsx", sheets=[sheet])
    assert str(cell(disk / "sortie/t.xlsx", "T", "A2").value) == "2026-10-25 13:50:00"
    # Sans fuseau, l'heure est prise telle quelle.
    assert str(cell(disk / "sortie/t.xlsx", "T", "A3").value) == "2026-10-25 09:00:00"
    assert "en heure de America/New_York" in tool.spec.description
    assert tool.settings_lines()[-3] == "fuseau : America/New_York"


async def test_a_percent_given_as_a_whole_number_warns(tool: ExcelTool) -> None:
    done = await call(tool, path="sortie/p.xlsx", sheets=[_one("pourcentage", 15)])
    assert done["warnings"] == [
        "F!A2 : 15 affiché 1500% — un pourcentage se donne en fraction (0.15 pour 15 %)"
    ]


@pytest.mark.parametrize(
    ("column_type", "value", "said"),
    [
        ("entier", 1.5, "1.5 n'est pas un\\(e\\) entier"),
        ("date", "demain", "'demain' n'est pas un\\(e\\) date"),
        ("montant_eur", "beaucoup", "n'est pas un\\(e\\) montant eur"),
        ("annee", "26", "n'est pas un\\(e\\) annee"),
        ("booleen", "peut-être", "n'est pas un\\(e\\) booleen"),
    ],
)
async def test_a_value_that_does_not_fit_its_type(
    tool: ExcelTool, column_type: str, value: JsonValue, said: str
) -> None:
    await refused(tool, said, path="sortie/v.xlsx", sheets=[_one(column_type, value)])


async def test_structure_faults_are_listed_together(tool: ExcelTool) -> None:
    sheets: list[JsonValue] = [
        {"name": "A/B", "columns": [{"header": "x"}]},
        {"name": "Suivi", "columns": [{"header": "x"}, {"header": "x"}]},
        {"name": "suivi", "columns": [{"header": "x"}]},
        {
            "name": "H",
            "columns": [{"header": "x"}],
            "rows": [{"cells": ["a", "b"], "role": "hypothese"}],
            "totals": {"y": "somme"},
        },
    ]
    with pytest.raises(ToolError) as caught:
        await tool.invoke({"path": "sortie/s.xlsx", "sheets": sheets}, CONTEXT)
    said = caught.value.message if hasattr(caught.value, "message") else str(caught.value)
    for part in (
        "les caractères [ ] : * ? / \\ sont interdits",
        "en-tête donné deux fois : x",
        "onglet 'suivi' donné deux fois",
        "H!2 : 2 valeurs pour 1 colonnes",
        "une hypothèse porte sa source en 'note'",
        "total sur une colonne inconnue 'y'",
    ):
        assert part in said, part


def test_sheet_names() -> None:
    assert sheet_name_problem("x" * 32) == f"onglet {'x' * 32!r} : 31 caractères au plus (32)"
    assert sheet_name_problem("'a") is not None
    assert sheet_name_problem(" ") == "nom d'onglet vide"
    assert sheet_name_problem("Suivi 2026") is None
    assert (quoted("Suivi"), quoted("Mes Hyp"), quoted("A1"), quoted("l'an")) == (
        "Suivi",
        "'Mes Hyp'",
        "'A1'",
        "'l''an'",
    )


# --- Formules --------------------------------------------------------------------------

SHEETS: Final = {"suivi": "Suivi", "mes hypothèses": "Mes Hypothèses"}


def _formula(formula: str) -> tuple[str, list[str], list[str]]:
    problems = Problems()
    written = check_formula(formula, "Suivi!C5", SHEETS, problems)
    return written, problems.errors, problems.warnings


@pytest.mark.parametrize(
    ("formula", "said"),
    [
        ("=SOMME(A1:A3)", "fonction française SOMME dans '=SOMME(A1:A3)' — écrire SUM"),
        ("=IF(A1>0;1;0)", "séparateur « ; »"),
        ("=SUM(A1:A3", "parenthèses déséquilibrées"),
        ("=XLOOKUP(A1,B:B,C:C)", "XLOOKUP n'est pas permise (fonction à débordement)"),
        ('=HYPERLINK("x")', "fonction HYPERLINK non permise"),
        ("=Autre!A1", "onglet inconnu 'Autre'"),
        ("='Rien'!A1", "onglet inconnu 'Rien'"),
        ("=Hypothèses!B2", "onglet inconnu 'Hypothèses'"),
        ("=A1048577", "référence A1048577 hors de la feuille"),
        ("=XFE1", "référence XFE1 hors de la feuille"),
        ('="a', "guillemet non fermé"),
    ],
)
def test_formula_refused(formula: str, said: str) -> None:
    _, errors, _ = _formula(formula)
    assert any(said in error for error in errors), errors


def test_formula_accepted_prefixed_and_warned() -> None:
    assert _formula("=SUM(C2:C9)") == ("=SUM(C2:C9)", [], [])
    assert _formula("='Mes Hypothèses'!$B$2*Suivi!C2") == (
        "='Mes Hypothèses'!$B$2*Suivi!C2",
        [],
        [],
    )
    written, errors, _ = _formula('=TEXTJOIN(", ",TRUE,A1:A3)&" ; textjoin("')
    assert (written, errors) == ('=_xlfn.TEXTJOIN(", ",TRUE,A1:A3)&" ; textjoin("', [])
    assert _formula("=ROUND(A1,2)")[2] == []
    _, _, warnings = _formula("=B5*1.05")
    assert warnings == [
        "Suivi!C5 : constante 1.05 dans '=B5*1.05' — une hypothèse se met dans sa propre "
        "cellule, citée par la formule"
    ]


async def test_formulas_are_checked_in_the_workbook(tool: ExcelTool, disk: Path) -> None:
    sheet = {
        "name": "Calcul",
        "columns": [
            {"header": "Taux", "type": "pourcentage"},
            {"header": "CA", "type": "montant_eur"},
        ],
        "rows": [{"cells": [0.1, "=A2*'Suivi'!C2"]}, {"cells": [None, "=SOMME(B2)"]}],
    }
    await refused(
        tool, "Calcul!B3 : fonction française SOMME", path="sortie/f.xlsx", sheets=[SUIVI, sheet]
    )
    sheet["rows"] = [{"cells": [0.1, "=A2*Suivi!C2"]}]
    done = await call(tool, path="sortie/f.xlsx", sheets=[SUIVI, sheet])
    linked = cell(disk / "sortie/f.xlsx", "Calcul", "B2")
    assert linked.value == "=A2*Suivi!C2"
    assert rgb(linked) == "00008000"
    assert done["sheets"][1]["formulas"] == 1


# --- Recalcul --------------------------------------------------------------------------


async def test_recalc_refuses_a_formula_in_error(disk: Path, soffice: str) -> None:
    tool = make(
        disk,
        {"name": "sortie", "path": "sortie", "rights": ["write"]},
        recalc=True,
        soffice=soffice,
    )
    sheet: dict[str, Any] = {
        "name": "R",
        "columns": [{"header": "A", "type": "entier"}, {"header": "B", "type": "decimal"}],
        "rows": [{"cells": [0, "=1/A2"]}, {"cells": [2, "=10/A3"]}],
    }
    await refused(
        tool,
        r"Recalcul : 1 formule\(s\) en erreur sur 2 — #DIV/0! : R!B2",
        path="sortie/r.xlsx",
        sheets=[sheet],
    )
    assert not (disk / "sortie" / "r.xlsx").exists()
    sheet["rows"] = [{"cells": [4, "=IFERROR(1/A2,0)"]}]
    done = await call(tool, path="sortie/r.xlsx", sheets=[sheet])
    assert done["recalc"] == {"formulas": 1, "errors": 0}
    # Le fichier livré reste celui d'openpyxl : ses formats, pas ceux de LibreOffice.
    assert cell(disk / "sortie/r.xlsx", "R", "A2").number_format == "#,##0"


async def test_recalc_timeout_and_failure(disk: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    fake = disk / "soffice"
    fake.write_text("#!/bin/sh\nexit 3\n", encoding="utf-8")
    fake.chmod(0o755)
    tool = make(
        disk,
        {"name": "sortie", "path": "sortie", "rights": ["write"]},
        recalc=True,
        soffice=str(fake),
        recalc_timeout=5,
    )
    sheet: JsonValue = {"name": "R", "columns": [{"header": "A"}], "rows": [{"cells": ["=1+1"]}]}
    await refused(tool, "Recalcul impossible", path="sortie/r.xlsx", sheets=[sheet])

    def slow(*args: Any, **kwargs: Any) -> None:
        raise subprocess.TimeoutExpired(cmd="soffice", timeout=5)

    monkeypatch.setattr(recalc_module.subprocess, "run", slow)
    await refused(tool, "Recalcul interrompu après 5 s", path="sortie/r.xlsx", sheets=[sheet])
    assert tool.settings_lines()[-2] == f"recalcul : {fake} (au plus 5 s)"


# --- Montage par la config, run, validate ----------------------------------------------


def excel_agent(**reference: Any) -> dict[str, Any]:
    return demo_agent(tools=[{"python": "loom.excel", **reference}])


def test_delete_in_approval_is_refused_when_the_config_loads(demo: ConfigFactory) -> None:
    """Par la référence de l'agent aussi, avant les règles communes des dossiers."""
    folders = [{"name": "s", "path": "/tmp", "rights": ["write"], "approval": {"delete": "never"}}]
    with pytest.raises(
        ConfigError, match=r"dossier s : le droit 'delete' ne sert pas à loom\.excel"
    ):
        load_config(demo(agents=[excel_agent(folders=folders)]))


def test_mounted_from_the_config(demo: ConfigFactory, tmp_path: Path) -> None:
    (tmp_path / "sortie").mkdir()
    folders = [{"name": "sortie", "path": "sortie", "rights": ["write"]}]
    path = demo(agents=[excel_agent(timeout=30, folders=folders)], storage=JSONL)
    agent = build_agent(load_config(path), "demo", InMemoryEventStore())
    (spec,) = agent.context.tools.specs
    assert (spec.name, spec.timeout, spec.approval) == ("excel", 30, "never")


def test_a_missing_extra_is_said_at_mount(
    demo: ConfigFactory, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    (tmp_path / "sortie").mkdir()
    # Le paquet et ses modules, déjà importés par ce fichier : à réimporter, sans openpyxl.
    package = "loom_ia.tools_builtin.excel"
    for name in [n for n in sys.modules if n == package or n.startswith(f"{package}.")]:
        monkeypatch.delitem(sys.modules, name)
    monkeypatch.setitem(sys.modules, "openpyxl", None)
    folders = [{"name": "sortie", "path": "sortie", "rights": ["write"]}]
    path = demo(agents=[excel_agent(folders=folders)], storage=JSONL)
    with pytest.raises(
        ConfigError,
        match=r"Outil fourni 'loom.excel' : dépendance 'openpyxl' absente "
        r"\(installer l'extra : loom-ia\[office\]\)",
    ):
        build_agent(load_config(path), "demo", InMemoryEventStore())


def test_common_settings_it_refuses(demo: ConfigFactory, tmp_path: Path) -> None:
    folders = [{"name": "sortie", "path": "sortie", "rights": ["write"]}]
    path = demo(agents=[excel_agent(approval="always", folders=folders)], storage=JSONL)
    with pytest.raises(ConfigError, match="approval ne se règle pas sur cet outil"):
        load_config(path)


async def test_a_run_writes_the_workbook(demo: ConfigFactory, tmp_path: Path) -> None:
    (tmp_path / "sortie").mkdir()
    script: list[dict[str, Any]] = [
        {
            "text": "J'écris le classeur.",
            "tool_calls": [
                {"name": "excel", "arguments": {"path": "sortie/r.xlsx", "sheets": [SUIVI]}}
            ],
        },
        {"text": "Classeur écrit."},
    ]
    folders = [{"name": "sortie", "path": "sortie", "rights": ["write"]}]
    path = demo(
        agents=[excel_agent(folders=folders)],
        models=[{**MODEL, "params": {"script": script}}],
        storage=JSONL,
    )
    store = InMemoryEventStore()
    agent = build_agent(load_config(path), "demo", store)
    try:
        run = await begin_run(agent.context, "Fais le suivi.")
        state = await drive(agent.context, run.run_id)
    finally:
        await agent.aclose()
    assert state.status is RunStatus.COMPLETED
    events = await store.read(DEFAULT_TENANT, run.session_id)
    [done] = [e.payload for e in events if isinstance(e.payload, ToolCompleted)]
    assert isinstance(done.output.data, dict)
    assert done.output.data["path"] == "sortie/r.xlsx"
    assert (tmp_path / "sortie" / "r.xlsx").is_file()


def test_validate_shows_the_folders_and_the_recalc(
    demo: ConfigFactory, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    (tmp_path / "sortie").mkdir()
    folders = [
        {"name": "sortie", "path": "sortie", "rights": ["write"], "approval": {"write": "always"}}
    ]
    path = demo(agents=[excel_agent(folders=folders)], storage=JSONL)
    assert main(["--config", str(path), "validate"]) == 0
    out = capsys.readouterr().out
    assert "Fournis    : loom.excel\n" in out
    assert (
        f"    outil excel : sortie → {(tmp_path / 'sortie').resolve()}\n"
        "    outil excel :   write : accordé, avec approbation\n"
    ) in out
    assert "    outil excel : recalcul : non\n" in out
