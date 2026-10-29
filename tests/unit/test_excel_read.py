# SPDX-License-Identifier: Apache-2.0
"""``loom.excel``, action ``read`` : lire un classeur, et un onglet .xlsx comme source (prototype).

Les classeurs d'essai sont fabriqués avec openpyxl, puis retouchés dans leur
zip pour les cas qu'openpyxl n'écrit pas : dimension fausse, entités XML,
bombe de décompression.
"""

import datetime as dt
import io
import zipfile
from collections.abc import Callable
from pathlib import Path
from typing import Any, Final

import pytest

pytest.importorskip("openpyxl", reason="extra 'office' absent")

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.formula import ArrayFormula

from loom_ia.core.model import DEFAULT_TENANT, RunId, SessionId
from loom_ia.core.ports import ToolContext, ToolError
from loom_ia.tools_builtin import BuiltinContext
from loom_ia.tools_builtin.excel import ExcelTool, build
from loom_ia.tools_builtin.excel import reader as reader_module
from loom_ia.tools_builtin.excel_params import ExcelParams

CONTEXT: Final = ToolContext(
    tenant_id=DEFAULT_TENANT,
    session_id=SessionId("s"),
    run_id=RunId("r"),
    call_id="c1",
    agent="demo",
)
SHEET: Final = "xl/worksheets/sheet1.xml"


@pytest.fixture
def disk(tmp_path: Path) -> Path:
    for name in ("donnees", "sortie", "valide"):
        (tmp_path / name).mkdir()
    return tmp_path


def make(disk: Path, *folders: dict[str, Any]) -> ExcelTool:
    model = ExcelParams.model_validate({"folders": list(folders)})
    return build(model, BuiltinContext(base_dir=disk))


@pytest.fixture
def tool(disk: Path) -> ExcelTool:
    return make(
        disk,
        {"name": "donnees", "path": "donnees", "rights": ["read"]},
        {"name": "sortie", "path": "sortie", "rights": ["write", "overwrite", "read"]},
        {"name": "valide", "path": "valide", "rights": ["read"], "approval": {"read": "always"}},
    )


async def call(tool: ExcelTool, **arguments: Any) -> dict[str, Any]:
    output = await tool.invoke(dict(arguments), CONTEXT)
    assert isinstance(output.data, dict)
    return output.data


async def refused(tool: ExcelTool, match: str, **arguments: Any) -> None:
    with pytest.raises(ToolError, match=match):
        await tool.invoke(dict(arguments), CONTEXT)


def sample(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Données"
    ws.append(["Client", "Montant", "Date", "Total"])
    for i in range(1, 5):
        ws.append([f"C-{i:03}", 100 * i, dt.datetime(2026, 10, i, 9, 30), f"=B{i + 1}*1.2"])
    ws["F2"] = dt.date(2026, 10, 28)
    ws["G2"] = ArrayFormula("G2:G3", "=B2:B3*2")
    ws["H2"] = dt.time(13, 0)
    ws["I2"] = dt.timedelta(hours=2, minutes=30)
    ws["J2"] = "#DIV/0!"
    ws["J2"].data_type = "e"
    ws["K2"] = "x" * 1_500
    hidden = wb.create_sheet("Caché")
    hidden["A1"] = "secret"
    hidden.sheet_state = "hidden"
    chart = wb.create_chartsheet("Graphique")
    bars = BarChart()
    reference = Reference(ws, min_col=2, min_row=1, max_row=5)
    bars.add_data(reference, titles_from_data=True)  # pyright: ignore[reportUnknownMemberType]
    chart.add_chart(bars)  # pyright: ignore[reportUnknownMemberType]
    wb.save(path)
    return path


def rewrite(source: Path, target: Path, edit: Callable[[str, bytes], bytes]) -> Path:
    archive = zipfile.ZipFile(source)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as written:
        for item in archive.infolist():
            written.writestr(item, edit(item.filename, archive.read(item.filename)))
    target.write_bytes(out.getvalue())
    return target


def lines(path: Path, count: int) -> Path:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["N"])
    for n in range(count):
        ws.append([n])
    wb.save(path)
    return path


# --- Arguments ----------------------------------------------------------------------


@pytest.mark.parametrize(
    ("arguments", "message"),
    [
        (
            {
                "action": "read",
                "path": "donnees/a.xlsx",
                "sheets": [{"name": "S", "columns": [{"header": "A"}]}],
            },
            "sheets ne sert pas à read",
        ),
        ({"path": "sortie/a.xlsx", "sheet": "S"}, "sheet ne sert pas à write"),
        ({"path": "sortie/a.xlsx"}, "write attend 'sheets'"),
        ({"action": "read", "path": "donnees/a.xlsx", "range": "A1"}, "'range' demande 'sheet'"),
        ({"action": "read", "path": "donnees/a.xlsx", "max_rows": 5000}, "max_rows"),
    ],
)
async def test_each_action_takes_its_own_fields(
    tool: ExcelTool, arguments: dict[str, Any], message: str
) -> None:
    await refused(tool, message, **arguments)


async def test_overwrite_false_on_a_read_is_not_a_fault(tool: ExcelTool, disk: Path) -> None:
    sample(disk / "donnees" / "a.xlsx")
    found = await call(tool, action="read", path="donnees/a.xlsx", overwrite=False)
    assert [s["name"] for s in found["sheets"]] == ["Données", "Caché"]


def test_the_description_and_schema_speak_of_reading(tool: ExcelTool) -> None:
    assert "action=read" in tool.spec.description
    assert "Lecture (action=read, dans : donnees, sortie, valide)" in tool.spec.description
    assert "'action'" in str(tool.spec.input_schema) or '"action"' in str(tool.spec.input_schema)


def test_without_any_read_right_the_description_says_so(disk: Path) -> None:
    only = make(disk, {"name": "sortie", "path": "sortie", "rights": ["write"]})
    assert "aucun dossier n'a le droit read" in only.spec.description


# --- Droits et approbation ------------------------------------------------------------


async def test_reading_needs_the_read_right(disk: Path) -> None:
    only = make(disk, {"name": "sortie", "path": "sortie", "rights": ["write"]})
    sample(disk / "sortie" / "a.xlsx")
    with pytest.raises(ToolError, match="read"):
        await only.invoke({"action": "read", "path": "sortie/a.xlsx"}, CONTEXT)


def test_a_read_has_no_effect_and_asks_only_where_the_folder_asks(tool: ExcelTool) -> None:
    free = tool.assess_call({"action": "read", "path": "donnees/a.xlsx"})
    assert (free.side_effects, free.approval, free.idempotent) == ("none", "never", True)
    asked = tool.assess_call({"action": "read", "path": "valide/a.xlsx"})
    assert asked.approval == "always" and asked.side_effects == "none"
    assert asked.reason == "lire le classeur valide/a.xlsx : approbation demandée pour read"


@pytest.mark.parametrize(
    ("path", "message"),
    [
        ("donnees/a.xls", r"format Excel 97-2003 \(\.xls\) non lu"),
        ("donnees/a.ods", r"OpenDocument \(\.ods\)"),
        ("donnees/a.csv", r"seuls les classeurs \.xlsx et \.xlsm sont lus"),
        ("donnees/absent.xlsx", "fichier introuvable"),
    ],
)
async def test_what_is_not_a_readable_workbook(tool: ExcelTool, path: str, message: str) -> None:
    await refused(tool, message, action="read", path=path)


# --- Ce qui est rendu -----------------------------------------------------------------


async def test_the_overview_lists_sheets_states_and_charts(tool: ExcelTool, disk: Path) -> None:
    sample(disk / "donnees" / "a.xlsx")
    found = await call(tool, action="read", path="donnees/a.xlsx")
    first, hidden = found["sheets"]
    assert (first["name"], first["state"], first["dimension"]) == ("Données", "visible", "A1:K5")
    assert first["preview"][0][:4] == ["Client", "Montant", "Date", "Total"]
    assert len(first["preview"]) == 5
    assert (hidden["name"], hidden["state"]) == ("Caché", "hidden")
    assert found["chartsheets"] == ["Graphique"]


async def test_cells_come_back_as_json(tool: ExcelTool, disk: Path) -> None:
    sample(disk / "donnees" / "a.xlsx")
    found = await call(tool, action="read", path="donnees/a.xlsx", sheet="Données")
    assert found["start"] == "A1" and found["row_count"] == 5 and not found["truncated"]
    header, first = found["rows"][0], found["rows"][1]
    # Lignes de même largeur : celle de la colonne non vide la plus à droite (K).
    assert header[:4] == ["Client", "Montant", "Date", "Total"] and len(header) == 11
    assert first[:3] == ["C-001", 100, "2026-10-01T09:30:00"]
    assert first[5] == "2026-10-28"  # une date sans heure
    assert first[7:10] == ["13:00:00", "2:30:00", "#DIV/0!"]
    assert first[10].endswith("…") and len(first[10]) == 1_001
    # Les formules, par adresse ; openpyxl ne garde pas de valeur calculée.
    assert found["formulas"]["D2"] == "=B2*1.2"
    assert found["formulas"]["G2"] == "{=B2:B3*2} sur G2:G3"
    assert first[3] is None
    assert any("sans valeur calculée" in note for note in found["notes"])
    assert any("1 texte(s) coupé(s)" in note for note in found["notes"])


async def test_a_range_says_where_it_starts(tool: ExcelTool, disk: Path) -> None:
    sample(disk / "donnees" / "a.xlsx")
    found = await call(tool, action="read", path="donnees/a.xlsx", sheet="Données", range="B3:C4")
    assert found["start"] == "B3"
    assert found["rows"] == [[200, "2026-10-02T09:30:00"], [300, "2026-10-03T09:30:00"]]
    await refused(
        tool, "plage illisible", action="read", path="donnees/a.xlsx", sheet="Données", range="1A:B"
    )


async def test_a_long_sheet_is_cut_and_says_how_to_go_on(tool: ExcelTool, disk: Path) -> None:
    lines(disk / "donnees" / "long.xlsx", 300)
    found = await call(tool, action="read", path="donnees/long.xlsx", sheet="Sheet")
    assert found["row_count"] == 200 and found["truncated"]
    assert found["next_start"] == "A201"
    assert any(
        "Lecture coupée par 200 lignes : la suite commence en A201" in n for n in found["notes"]
    )
    rest = await call(
        tool, action="read", path="donnees/long.xlsx", sheet="Sheet", range="A201:A400"
    )
    assert rest["row_count"] == 101 and not rest["truncated"]


async def test_data_after_empty_rows_still_counts_as_more(tool: ExcelTool, disk: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws["A1"], ws["A2"], ws["A50"] = "a", "b", "loin"
    wb.save(disk / "donnees" / "trou.xlsx")
    found = await call(tool, action="read", path="donnees/trou.xlsx", sheet="Sheet", max_rows=10)
    assert found["rows"] == [["a"], ["b"]] and found["truncated"]


async def test_a_hidden_sheet_says_it(tool: ExcelTool, disk: Path) -> None:
    sample(disk / "donnees" / "a.xlsx")
    found = await call(tool, action="read", path="donnees/a.xlsx", sheet="Caché")
    assert found["rows"] == [["secret"]] and found["state"] == "hidden"


@pytest.mark.parametrize(
    ("sheet", "message"),
    [
        ("Absent", r"onglet 'Absent' introuvable \(onglets : Données, Caché"),
        ("Graphique", "onglet graphique"),
    ],
)
async def test_unknown_or_chart_sheets(
    tool: ExcelTool, disk: Path, sheet: str, message: str
) -> None:
    sample(disk / "donnees" / "a.xlsx")
    await refused(tool, message, action="read", path="donnees/a.xlsx", sheet=sheet)


async def test_a_lying_dimension_does_not_cut_the_sheet(tool: ExcelTool, disk: Path) -> None:
    """Le fichier annonce « A1 » : sans reset_dimensions, la lecture seule ne rendait que A1."""
    source = sample(disk / "donnees" / "a.xlsx")
    rewrite(
        source,
        disk / "donnees" / "faux.xlsx",
        lambda name, data: (
            data.replace(b'<dimension ref="A1:K5" />', b'<dimension ref="A1" />')
            if name == SHEET
            else data
        ),
    )
    found = await call(tool, action="read", path="donnees/faux.xlsx", sheet="Données")
    assert found["row_count"] == 5


async def test_a_workbook_written_by_excel_reads_back(tool: ExcelTool, disk: Path) -> None:
    sheets = [
        {
            "name": "Ventes",
            "columns": [{"header": "Client"}, {"header": "HT", "type": "montant_eur"}],
            "rows": [{"cells": ["C-001", 10]}, {"cells": ["C-002", 20]}],
            "totals": {"HT": "somme"},
        }
    ]
    await call(tool, path="sortie/v.xlsx", sheets=sheets)
    found = await call(tool, action="read", path="sortie/v.xlsx", sheet="Ventes")
    assert found["rows"][:3] == [["Client", "HT"], ["C-001", 10], ["C-002", 20]]
    assert found["formulas"] == {"B4": "=SUM(B2:B3)"}
    assert any("1 formule(s) sans valeur calculée" in note for note in found["notes"])


# --- Un .xlsx comme source --------------------------------------------------------------


async def test_a_sheet_feeds_a_new_workbook(tool: ExcelTool, disk: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Données"
    ws.append(["Client", "Montant", "Date"])
    for i in range(1, 5):
        ws.append([f"C-{i:03}", 100 * i, dt.datetime(2026, 10, i, 9, 30)])
    wb.save(disk / "donnees" / "a.xlsx")
    sheets = [
        {
            "name": "Copie",
            "columns": [{"header": "Client"}, {"header": "Date", "type": "date_heure"}],
            "data": {"path": "donnees/a.xlsx", "sheet": "Données"},
        }
    ]
    written = await call(tool, path="sortie/copie.xlsx", sheets=sheets)
    assert written["sheets"] == [{"name": "Copie", "rows": 4, "formulas": 0}]
    await refused(
        tool,
        r"'sheet' ne sert qu'à un \.xlsx",
        path="sortie/x.xlsx",
        sheets=[{**sheets[0], "data": {"path": "donnees/b.csv", "sheet": "S"}}],
    )


# --- Sécurité -------------------------------------------------------------------------------

LOL: Final = (
    b'<!DOCTYPE w [<!ENTITY a "aaaaaaaaaa">'
    + b"".join(
        b'<!ENTITY %s "%s">'
        % (
            bytes(f"e{n}", "ascii"),
            b"".join([bytes(f"&e{n - 1};" if n > 1 else "&a;", "ascii")] * 10),
        )
        for n in range(1, 10)
    )
    + b"]>"
)


@pytest.mark.parametrize(
    ("prologue", "entity", "message"),
    [
        (LOL, b"&e9;", "XML refusé"),
        (b'<!DOCTYPE w [<!ENTITY x SYSTEM "file:///etc/hostname">]>', b"&x;", "XML refusé"),
    ],
)
async def test_xml_entities_are_refused(
    tool: ExcelTool, disk: Path, prologue: bytes, entity: bytes, message: str
) -> None:
    """« Billion laughs » : expat refuse l'amplification ; entité externe : jamais chargée."""
    source = sample(disk / "donnees" / "a.xlsx")
    rewrite(
        source,
        disk / "donnees" / "x.xlsx",
        lambda name, data: (
            prologue + data.replace(b"<t>Client</t>", b"<t>" + entity + b"</t>")
            if name == SHEET
            else data
        ),
    )
    await refused(tool, message, action="read", path="donnees/x.xlsx", sheet="Données")


async def test_a_decompression_bomb_is_refused_before_openpyxl(
    tool: ExcelTool, disk: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = sample(disk / "donnees" / "a.xlsx")

    def bomb(name: str, data: bytes) -> bytes:
        if name != SHEET:
            return data
        head, tail = data.split(b"<sheetData>")[0], data.split(b"</sheetData>")[1]
        body = b"<row><c><v>0</v></c></row>" * 400_000
        return head + b"<sheetData>" + body + b"</sheetData>" + tail

    rewrite(source, disk / "donnees" / "bombe.xlsx", bomb)
    # 10 Mo de XML pour un fichier de quelques dizaines de Ko : la borne, ramenée à 5 Mo.
    assert (disk / "donnees" / "bombe.xlsx").stat().st_size < 100_000
    monkeypatch.setattr(reader_module, "MAX_UNCOMPRESSED", 5_000_000)
    await refused(tool, "une fois décompressé", action="read", path="donnees/bombe.xlsx")


async def test_the_uncompressed_size_is_bounded(
    tool: ExcelTool, disk: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    sample(disk / "donnees" / "a.xlsx")
    monkeypatch.setattr(reader_module, "MAX_UNCOMPRESSED", 1_000)
    await refused(
        tool, "une fois décompressé, au-delà de 1000", action="read", path="donnees/a.xlsx"
    )


@pytest.mark.parametrize(
    ("content", "message"), [(b"", "pas un classeur .xlsx lisible"), (b"a;b\n1;2\n", "renommé")]
)
async def test_a_file_that_is_not_a_workbook(
    tool: ExcelTool, disk: Path, content: bytes, message: str
) -> None:
    (disk / "donnees" / "faux.xlsx").write_bytes(content)
    await refused(tool, message, action="read", path="donnees/faux.xlsx")


async def test_a_zip_that_is_not_a_workbook(tool: ExcelTool, disk: Path) -> None:
    with zipfile.ZipFile(disk / "donnees" / "z.xlsx", "w") as archive:
        archive.writestr("a.txt", "rien")
    await refused(tool, r"xl/workbook\.xml absent", action="read", path="donnees/z.xlsx")


async def test_a_range_after_empty_rows_lands_on_the_right_row(tool: ExcelTool, disk: Path) -> None:
    """La lecture part de la ligne 1 et compte elle-même : les lignes absentes comptent aussi."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws["A1"], ws["A50"], ws["B60"] = "haut", "loin", "=1+1"
    wb.save(disk / "donnees" / "trou.xlsx")
    found = await call(
        tool, action="read", path="donnees/trou.xlsx", sheet="Sheet", range="A50:B60"
    )
    # La colonne B garde sa formule sans valeur : les lignes ont deux colonnes.
    assert found["start"] == "A50" and found["rows"][0] == ["loin", None]
    assert found["formulas"] == {"B60": "=1+1"}


async def test_a_long_read_stops_by_itself(
    tool: ExcelTool, disk: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Un fil ne s'annule pas : sans échéance, il lirait encore après le délai de l'outil."""
    lines(disk / "donnees" / "long.xlsx", 3_000)
    monkeypatch.setattr(reader_module, "READ_DEADLINE", -1.0)
    await refused(
        tool,
        "lecture arrêtée après",
        action="read",
        path="donnees/long.xlsx",
        sheet="Sheet",
        range="A2500:A2600",
    )


# --- Relecture indépendante : un test par défaut trouvé -------------------------------------


def edited(source: Path, target: Path, old: bytes, new: bytes) -> Path:
    return rewrite(
        source, target, lambda name, data: data.replace(old, new) if name == SHEET else data
    )


@pytest.mark.parametrize(
    ("old", "new"),
    [
        # Un nombre qui n'en est pas un.
        (b'<c r="B2" t="n"><v>100</v></c>', b'<c r="B2" t="n"><v>abc</v></c>'),
        # Une date ISO invalide.
        (b'<c r="B2" t="n"><v>100</v></c>', b'<c r="B2" t="d"><v>2026-13-45</v></c>'),
    ],
)
async def test_a_broken_cell_is_a_tool_error_not_an_internal_one(
    tool: ExcelTool, disk: Path, old: bytes, new: bytes
) -> None:
    source = sample(disk / "donnees" / "a.xlsx")
    assert old in zipfile.ZipFile(source).read(SHEET)
    edited(source, disk / "donnees" / "x.xlsx", old, new)
    for arguments in ({}, {"sheet": "Données"}):
        await refused(tool, "classeur illisible", action="read", path="donnees/x.xlsx", **arguments)


async def test_a_chart_sheet_as_a_source_is_said(tool: ExcelTool, disk: Path) -> None:
    sample(disk / "donnees" / "a.xlsx")
    sheets = [
        {
            "name": "S",
            "columns": [{"header": "Client"}],
            "data": {"path": "donnees/a.xlsx", "sheet": "Graphique"},
        }
    ]
    await refused(tool, "onglet graphique", path="sortie/s.xlsx", sheets=sheets)


async def test_the_result_has_a_size_budget(tool: ExcelTool, disk: Path) -> None:
    """1 000 lignes de longs textes : la lecture s'arrête au budget et le dit."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    for _ in range(1_000):
        ws.append(["x" * 900] * 20)
    wb.save(disk / "donnees" / "large.xlsx")
    found = await call(
        tool, action="read", path="donnees/large.xlsx", sheet="Sheet", max_rows=1_000
    )
    assert len(str(found["rows"])) < 60_000
    assert found["truncated"] and found["next_start"] == f"A{found['row_count'] + 1}"
    assert any("la taille du résultat (40000 caractères)" in note for note in found["notes"])


async def test_columns_beyond_the_bound_are_said(tool: ExcelTool, disk: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.cell(1, 1, "a")
    ws.cell(2, 250, "loin à droite")
    wb.save(disk / "donnees" / "large.xlsx")
    found = await call(tool, action="read", path="donnees/large.xlsx", sheet="Sheet")
    assert any("Colonnes coupées après GR (200 au plus)" in note for note in found["notes"])
    right = await call(
        tool, action="read", path="donnees/large.xlsx", sheet="Sheet", range="IP1:IP5"
    )
    assert right["rows"] == [[None], ["loin à droite"]] and right["start"] == "IP1"


async def test_a_source_keeps_long_texts_and_refuses_uncomputed_formulas(
    tool: ExcelTool, disk: Path
) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["Libellé", "TTC"])
    ws.append(["y" * 1_500, 12])
    wb.save(disk / "donnees" / "texte.xlsx")
    sheet = {"name": "S", "columns": [{"header": "Libellé"}, {"header": "TTC", "type": "decimal"}]}
    done = await call(
        tool, path="sortie/s.xlsx", sheets=[{**sheet, "data": {"path": "donnees/texte.xlsx"}}]
    )
    assert done["sheets"][0]["rows"] == 1
    ws["B2"] = "=10*1.2"
    wb.save(disk / "donnees" / "formule.xlsx")
    await refused(
        tool,
        r"1 formule\(s\) sans valeur calculée dans la source",
        path="sortie/t.xlsx",
        sheets=[{**sheet, "data": {"path": "donnees/formule.xlsx"}}],
    )
    ws["B1"] = "Libellé"
    wb.save(disk / "donnees" / "double.xlsx")
    await refused(
        tool,
        "en-têtes en double : Libellé",
        path="sortie/u.xlsx",
        sheets=[{**sheet, "data": {"path": "donnees/double.xlsx"}}],
    )


@pytest.mark.parametrize(
    ("given", "start", "message"),
    [
        ("C3:B2", "B2", None),  # dans l'ordre, comme Excel
        ("A0", None, "une ligne va de 1 à 1048576"),
        ("XFE1", None, "une colonne va de A à XFD"),
    ],
)
async def test_ranges_are_put_in_order_or_refused(
    tool: ExcelTool, disk: Path, given: str, start: str | None, message: str | None
) -> None:
    sample(disk / "donnees" / "a.xlsx")
    arguments = {"action": "read", "path": "donnees/a.xlsx", "sheet": "Données", "range": given}
    if message is not None:
        await refused(tool, message, **arguments)
        return
    found = await call(tool, **arguments)
    assert found["start"] == start and found["rows"][0] == [100, "2026-10-01T09:30:00"]


async def test_a_formula_without_value_stays_in_the_rows(tool: ExcelTool, disk: Path) -> None:
    sample(disk / "donnees" / "a.xlsx")
    found = await call(tool, action="read", path="donnees/a.xlsx", sheet="Données", range="D5")
    assert found["rows"] == [[None]] and found["row_count"] == 1
    assert found["formulas"] == {"D5": "=B5*1.2"}


async def test_the_next_start_counts_the_rows_walked(tool: ExcelTool, disk: Path) -> None:
    """Coupé après des lignes vides : la suite part de la ligne qui suit les lignes parcourues."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws["A1"], ws["A2"], ws["A3"], ws["A9"] = "a", "b", "c", "suite"
    wb.save(disk / "donnees" / "trous.xlsx")
    found = await call(tool, action="read", path="donnees/trous.xlsx", sheet="Sheet", max_rows=5)
    assert found["rows"] == [["a"], ["b"], ["c"]] and found["next_start"] == "A6"


async def test_text_that_looks_like_a_formula_and_infinity(tool: ExcelTool, disk: Path) -> None:
    source = sample(disk / "donnees" / "a.xlsx")
    data = zipfile.ZipFile(source).read(SHEET)
    assert b'<c r="B3" t="n"><v>200</v></c>' in data
    edited(
        source,
        disk / "donnees" / "b.xlsx",
        b'<c r="B3" t="n"><v>200</v></c>',
        b'<c r="B3" t="inlineStr"><is><t>=pas une formule</t></is></c>',
    )
    found = await call(tool, action="read", path="donnees/b.xlsx", sheet="Données", range="B3")
    assert found["rows"] == [["=pas une formule"]] and "formulas" not in found
    edited(
        source,
        disk / "donnees" / "c.xlsx",
        b'<c r="B3" t="n"><v>200</v></c>',
        b'<c r="B3" t="n"><v>1E+999</v></c>',
    )
    found = await call(tool, action="read", path="donnees/c.xlsx", sheet="Données", range="B3")
    assert found["rows"] == [["inf"]]


# --- Décisions du 27/10 -------------------------------------------------------------------


async def test_an_agent_may_only_read(disk: Path) -> None:
    """D-4 : un dossier read suffit ; l'outil le dit, et refuse d'écrire."""
    reader = make(disk, {"name": "donnees", "path": "donnees", "rights": ["read"]})
    assert "Aucun dossier n'a le droit write : l'outil ne fait que lire" in reader.spec.description
    assert dict(reader.settings().options)["lecture des classeurs"] == "oui"
    sample(disk / "donnees" / "a.xlsx")
    found = await call(reader, action="read", path="donnees/a.xlsx")
    assert found["sheets"][0]["name"] == "Données"
    sheets = [{"name": "S", "columns": [{"header": "A"}]}]
    await refused(reader, "write", path="donnees/b.xlsx", sheets=sheets)


def test_a_writer_without_read_says_it_cannot_read(disk: Path) -> None:
    writer = make(disk, {"name": "sortie", "path": "sortie", "rights": ["write"]})
    assert dict(writer.settings().options)["lecture des classeurs"] == "aucun dossier read"


@pytest.mark.parametrize(
    ("expat", "lxml", "said"),
    [
        ((2, 2, 10), None, r"expat 2\.2\.10 : 2\.4\.0 ou plus attendu"),
        (
            (2, 7, 1),
            ((4, 9, 4, 0), (2, 10, 3)),
            r"lxml 4\.9\.4 / libxml2 2\.10\.3 .* OPENPYXL_LXML",
        ),
        ((2, 7, 1), ((6, 1, 3, 0), (2, 14, 6)), None),
    ],
)
async def test_reading_is_refused_with_unsafe_xml_parsers(
    disk: Path,
    monkeypatch: pytest.MonkeyPatch,
    expat: tuple[int, int, int],
    lxml: tuple[tuple[int, ...], tuple[int, ...]] | None,
    said: str | None,
) -> None:
    """D-6 : expat (les feuilles) et, si openpyxl s'en sert, lxml (workbook.xml, styles)."""
    import pyexpat
    import sys
    import types

    import openpyxl.xml

    monkeypatch.setattr(pyexpat, "version_info", expat)
    monkeypatch.setattr(openpyxl.xml, "LXML", lxml is not None)
    if lxml is not None:
        fake = types.ModuleType("lxml.etree")
        fake.LXML_VERSION, fake.LIBXML_VERSION = lxml  # type: ignore[attr-defined]
        monkeypatch.setitem(sys.modules, "lxml.etree", fake)
    assert (reader_module.xml_safety() is None) == (said is None)
    if said is None:
        return
    import re

    assert re.search(said, reader_module.xml_safety() or "")
    tool = make(disk, {"name": "donnees", "path": "donnees", "rights": ["read", "write"]})
    assert dict(tool.settings().options)["lecture des classeurs"].startswith("non (")
    assert "Lecture (action=read) : indisponible ici" in tool.spec.description
    sample(disk / "donnees" / "a.xlsx")
    with pytest.raises(ToolError, match="lecture des classeurs désactivée ici"):
        await tool.invoke({"action": "read", "path": "donnees/a.xlsx"}, CONTEXT)


def padded(source: Path, target: Path, member: str, *, head: bytes = b"", tail: int = 0) -> Path:
    """Le classeur, avec un membre précédé de ``head`` et suivi de ``tail`` espaces."""

    def edit(name: str, data: bytes) -> bytes:
        return head + data + b" " * tail if name == member else data

    return rewrite(source, target, edit)


@pytest.mark.parametrize(
    ("member", "head", "edit", "admitted"),
    [
        # D-2, relecture 2 : les styles sont lus en entier (46 Mo : 14 s, 782 Mo).
        ("xl/styles.xml", b"", None, False),
        # Une feuille qui annonce sa dimension est lue au fil de l'eau : admise.
        (SHEET, b"", None, True),
        # Sans dimension, elle serait parcourue à l'ouverture : refusée.
        (SHEET, b"", (b'<dimension ref="A1:K5" />', b""), False),
        # Un commentaire ne trompe pas le contrôle : c'est le parseur qui lit.
        (SHEET, b"", (b'<dimension ref="A1:K5" />', b'<!--<dimension ref="A1"/>-->'), False),
        # Un <x:dimension> préfixé, comme l'écrit l'OpenXML SDK, compte.
        (
            SHEET,
            b"",
            (
                b'<dimension ref="A1:K5" />',
                b'<x:dimension xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/'
                b'main" ref="A1:K5" />',
            ),
            True,
        ),
    ],
)
async def test_what_openpyxl_would_read_whole_is_bounded(
    tool: ExcelTool,
    disk: Path,
    monkeypatch: pytest.MonkeyPatch,
    member: str,
    head: bytes,
    edit: tuple[bytes, bytes] | None,
    admitted: bool,
) -> None:
    source = sample(disk / "donnees" / "a.xlsx")
    if edit is not None:
        source = edited(source, disk / "donnees" / "e.xlsx", *edit)
    padded(source, disk / "donnees" / "p.xlsx", member, head=head, tail=60_000)
    monkeypatch.setattr(reader_module, "MAX_WHOLE_MEMBER", 30_000)
    arguments = {"action": "read", "path": "donnees/p.xlsx", "sheet": "Données"}
    if admitted:
        assert (await call(tool, **arguments))["row_count"] == 5
        return
    await refused(tool, f"{member} fait .* au-delà de 30000, et serait lu en entier", **arguments)


async def test_a_broken_member_during_the_check_is_a_tool_error(
    tool: ExcelTool, disk: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Relecture 2 : zlib.error ou membre chiffré pendant le contrôle du zip."""
    import zlib

    def broken(archive: object, member: object) -> bool:
        raise zlib.error("Error -3 while decompressing data")

    sample(disk / "donnees" / "a.xlsx")
    monkeypatch.setattr(reader_module, "MAX_WHOLE_MEMBER", 0)
    monkeypatch.setattr(reader_module, "_sized_sheet", broken)
    await refused(
        tool, "classeur illisible \\(error: Error -3", action="read", path="donnees/a.xlsx"
    )


async def test_a_source_workbook_is_refused_when_reading_is_unsafe(
    disk: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import pyexpat

    monkeypatch.setattr(pyexpat, "version_info", (2, 2, 10))
    tool = make(
        disk,
        {"name": "donnees", "path": "donnees", "rights": ["read"]},
        {"name": "sortie", "path": "sortie", "rights": ["write"]},
    )
    sample(disk / "donnees" / "a.xlsx")
    sheets = [{"name": "S", "columns": [{"header": "Client"}], "data": {"path": "donnees/a.xlsx"}}]
    await refused(tool, "lecture des classeurs désactivée ici", path="sortie/s.xlsx", sheets=sheets)


async def test_a_row_that_does_not_fit_the_budget_starts_the_next_read(
    tool: ExcelTool, disk: Path
) -> None:
    """Relecture 2 : la ligne qui ferait dépasser le budget n'est pas ajoutée."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["court"])
    ws.append(["y" * 900] * 50)
    wb.save(disk / "donnees" / "b.xlsx")
    found = await call(tool, action="read", path="donnees/b.xlsx", sheet="Sheet")
    assert found["rows"] == [["court"]] and found["next_start"] == "A2"
    # Seule, une ligne plus longue que le budget est rendue entière.
    alone = await call(tool, action="read", path="donnees/b.xlsx", sheet="Sheet", range="A2")
    assert alone["row_count"] == 1


def test_a_tool_that_only_reads_has_no_effect(disk: Path) -> None:
    """Relecture 2 : sans dossier write, l'outil se déclare sans effet, et se rejoue en variante."""
    from loom_ia.core.ports import strictest_side_effects

    reader = make(disk, {"name": "donnees", "path": "donnees", "rights": ["read"]})
    assert reader.spec.side_effects == "none"
    assert strictest_side_effects(reader, {"action": "read", "path": "donnees/a.xlsx"}) == "none"
    assert "Types de colonne" not in reader.spec.description  # le mode d'emploi de l'écriture


async def test_local_files_points_to_excel_for_a_workbook(disk: Path) -> None:
    from loom_ia.tools_builtin.local_files import build as build_files
    from loom_ia.tools_builtin.local_files_params import LocalFilesParams

    files = build_files(
        LocalFilesParams.model_validate(
            {"folders": [{"name": "donnees", "path": "donnees", "rights": ["read"]}]}
        ),
        BuiltinContext(base_dir=disk),
    )
    sample(disk / "donnees" / "a.xlsx")
    with pytest.raises(ToolError, match=r"un classeur se lit avec l'outil excel \(action=read\)"):
        await files.invoke({"action": "read", "path": "donnees/a.xlsx"}, CONTEXT)
