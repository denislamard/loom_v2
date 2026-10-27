# SPDX-License-Identifier: Apache-2.0
"""Lecture d'un classeur .xlsx : ce qu'il contient, rendu en JSON pour le modèle.

Un classeur est un zip de XML qui vient de n'importe où. Avant openpyxl :

- le zip est contrôlé (taille décompressée, nombre de membres) : un fichier
  de 762 Ko qui se décompresse en 312 Mo bloquait openpyxl plus de deux
  minutes en mode normal ;
- openpyxl ne s'ouvre qu'en **lecture seule** (``read_only``), qui lit la
  feuille au fil de l'eau et s'arrête à la borne de lignes ; le mode normal
  charge tout en mémoire (456 Mo pour 100 000 lignes) ;
- les entités XML sont laissées au parseur de la bibliothèque standard
  (expat), qui refuse l'amplification (« billion laughs ») et ne charge pas
  d'entité externe.

Une passe en lecture seule pour les cellules et leurs formules ; une seconde,
seulement si la plage a des formules, pour les valeurs que le classeur a
gardées en cache. Un classeur écrit par openpyxl (donc par ``excel`` lui-même)
n'a pas de valeur en cache : ses formules sont rendues sans valeur, et le
résultat le dit.
"""

import datetime as dt
import importlib
import io
import json
import math
import time
import zipfile
from collections.abc import Generator, Iterator
from contextlib import contextmanager
from dataclasses import dataclass
from typing import Any, Final, cast
from xml.etree.ElementTree import Element, ParseError, XMLPullParser

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.formula import ArrayFormula
from pydantic import JsonValue

from loom_ia.core.ports import ToolError
from loom_ia.tools_builtin.excel.arguments import (
    MAX_COLUMNS,
    MAX_DATA_ROWS,
    MAX_SHEET_COLUMNS,
    MAX_SHEET_ROWS,
    MAX_SOURCE_BYTES,
)

# Un .xlsx courant se compresse 4 à 10 fois (mesuré : 3,9 à 9,8). La borne porte sur la
# taille décompressée, pas sur le taux : une colonne de textes identiques se compresse
# 163 fois sans être une bombe, alors qu'une bombe dépasse toujours la taille.
MAX_UNCOMPRESSED: Final = 10 * MAX_SOURCE_BYTES
# Ce que load_workbook lit en entier avant de rendre la main, hors de toute échéance :
# textes partagés, styles, classeur, et toute feuille qui n'annonce pas sa dimension.
# Mesuré : 119 Mo de textes partagés, 5,7 s et 269 Mo par ouverture ; 46 Mo de styles,
# 14,1 s et 782 Mo. Au-delà, un membre n'est admis que s'il est une feuille qui annonce
# sa dimension : openpyxl la lit alors au fil de l'eau, sous l'échéance.
MAX_WHOLE_MEMBER: Final = 20_000_000
# Où chercher <dimension> : il suit l'en-tête de la feuille, bien avant les données.
DIMENSION_WINDOW: Final = 65_536
MAIN_NS: Final = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
# Versions sûres face aux entités XML : expat refuse l'amplification depuis 2.4.0 ;
# libxml2 depuis 2.11 ; lxml ne résout plus que les entités internes depuis 5.0.
SAFE_EXPAT: Final = (2, 4, 0)
SAFE_LXML: Final = (5, 0, 0)
SAFE_LIBXML2: Final = (2, 11, 0)
MAX_MEMBERS: Final = 10_000
MAX_CELL_CHARS: Final = 1_000
PREVIEW_ROWS: Final = 5
PREVIEW_COLUMNS: Final = 20
PREVIEW_CHARS: Final = 100
# Ce que rend une lecture, en caractères JSON : sous le seuil de déport par défaut
# (50 000). Sans lui, 1 000 lignes de 200 textes de 1 000 caractères faisaient
# 200 millions de caractères pour un fichier de 423 Ko.
READ_CHARS: Final = 40_000
# Sous le délai par défaut d'un outil (30 s) : au-delà, la lecture s'arrête d'elle-même.
READ_DEADLINE: Final = 20.0
READABLE: Final = (".xlsx", ".xlsm")
# Formats qu'openpyxl ne lit pas : le dire plutôt que « pas un zip ».
OTHER_FORMATS: Final = {
    ".xls": "Excel 97-2003 (.xls)",
    ".xlsb": "Excel binaire (.xlsb)",
    ".ods": "OpenDocument (.ods)",
}


@dataclass(frozen=True, slots=True)
class Read:
    """Les cellules lues d'une plage : valeurs, formules, et ce qu'il faut en dire."""

    first_row: int
    first_column: int
    rows: list[list[JsonValue]]
    formulas: dict[str, str]
    without_value: int
    # Première ligne non rendue, s'il en reste une non vide plus bas.
    next_row: int | None
    # Dernière colonne rendue, si des colonnes plus à droite ont été laissées.
    columns_cut_after: str | None
    cut_cells: int
    budget_hit: bool


@contextmanager
def _guarded(where: str) -> Generator[None]:
    """Toute faute du fichier devient une ``ToolError`` pour le modèle, jamais une erreur interne.

    openpyxl lève ce que lève son parseur : ``ParseError`` (expat), mais aussi
    ``ValueError``, ``IndexError``, ``BadZipFile`` (CRC faux), ``AttributeError``
    sur un onglet graphique, et ``XMLSyntaxError`` quand lxml est installé.
    """
    try:
        yield
    except ToolError:
        raise
    except ParseError as exc:
        raise ToolError(f"{where} : XML refusé ({exc}).") from None
    except Exception as exc:
        raise ToolError(f"{where} : classeur illisible ({type(exc).__name__}: {exc}).") from None


def check_suffix(path: str) -> None:
    lowered = path.lower()
    for suffix, name in OTHER_FORMATS.items():
        if lowered.endswith(suffix):
            raise ToolError(
                f"{path} : format {name} non lu — l'enregistrer en .xlsx depuis Excel "
                "ou LibreOffice."
            )
    if not lowered.endswith(READABLE):
        raise ToolError(f"{path} : seuls les classeurs .xlsx et .xlsm sont lus.")


def check_archive(data: bytes, where: str) -> None:
    """Le zip avant openpyxl : un classeur, de taille raisonnable, qui s'ouvre vite.

    Sous ``_guarded`` : un membre abîmé (``zlib.error``, membre chiffré) devient
    une ``ToolError``, pas une erreur interne.
    """
    with _guarded(where):
        _check_archive(data, where)


def _check_archive(data: bytes, where: str) -> None:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        # Un classeur protégé par mot de passe est un conteneur OLE, pas un zip.
        raise ToolError(
            f"{where} : pas un classeur .xlsx lisible (fichier vide, renommé, ou protégé "
            "par mot de passe)."
        ) from None
    with archive:
        members = archive.infolist()
        if len(members) > MAX_MEMBERS:
            raise ToolError(f"{where} : {len(members)} fichiers dans le classeur, refusé.")
        total = sum(member.file_size for member in members)
        if total > MAX_UNCOMPRESSED:
            raise ToolError(
                f"{where} : {total} octets une fois décompressé, au-delà de {MAX_UNCOMPRESSED}."
            )
        if "xl/workbook.xml" not in {member.filename for member in members}:
            raise ToolError(f"{where} : pas un classeur Excel (xl/workbook.xml absent).")
        for member in members:
            if member.file_size > MAX_WHOLE_MEMBER and not _sized_sheet(archive, member):
                raise ToolError(
                    f"{where} : {member.filename} fait {member.file_size} octets décompressés, "
                    f"au-delà de {MAX_WHOLE_MEMBER}, et serait lu en entier à l'ouverture "
                    "(textes partagés, styles, ou feuille qui n'annonce pas sa taille) ; "
                    "l'exporter en CSV, ou l'enregistrer depuis Excel ou LibreOffice."
                )


def _sized_sheet(archive: zipfile.ZipFile, member: zipfile.ZipInfo) -> bool:
    """Vrai pour une feuille qui annonce sa dimension avant ses données.

    Lu par le parseur, pas par une recherche d'octets : un commentaire
    ``<!--<dimension-->`` ne compte pas, un ``<x:dimension>`` préfixé si. Le
    nom du membre ne compte pas : openpyxl suit les rels, pas les chemins.
    """
    parser: XMLPullParser[Element] = XMLPullParser(events=("start",))
    with archive.open(member) as stream:
        parser.feed(stream.read(DIMENSION_WINDOW))
    root: str | None = None
    events = cast(Iterator[tuple[str, Element]], parser.read_events())
    for _, found in events:
        tag = found.tag
        if root is None:
            root = tag
            if root != f"{MAIN_NS}worksheet":
                return False
        elif tag == f"{MAIN_NS}dimension":
            return True
        elif tag == f"{MAIN_NS}sheetData":
            return False
    return False


def xml_safety() -> str | None:
    """Pourquoi les classeurs ne peuvent pas être lus sans risque ici ; ``None`` s'ils le peuvent.

    Les feuilles passent par expat ; avec lxml installé, openpyxl lit aussi
    ``workbook.xml`` et les styles avec libxml2. Une version trop ancienne
    laisserait passer l'amplification d'entités (« billion laughs »).
    """
    import pyexpat

    from openpyxl.xml import LXML

    if pyexpat.version_info < SAFE_EXPAT:
        found = ".".join(map(str, pyexpat.version_info))
        return f"expat {found} : 2.4.0 ou plus attendu (amplification d'entités XML)"
    if LXML:
        # lxml n'est pas une dépendance de loom : lu sans l'importer statiquement.
        etree = importlib.import_module("lxml.etree")
        lxml_version = cast(tuple[int, ...], etree.LXML_VERSION)[:3]
        libxml2_version = cast(tuple[int, ...], etree.LIBXML_VERSION)
        if lxml_version < SAFE_LXML or libxml2_version < SAFE_LIBXML2:
            lxml = ".".join(map(str, lxml_version))
            libxml2 = ".".join(map(str, libxml2_version))
            return (
                f"lxml {lxml} / libxml2 {libxml2} : lxml 5.0 et libxml2 2.11 ou plus attendus "
                "(entités XML) ; mettre lxml à jour, ou poser OPENPYXL_LXML=False"
            )
    return None


def _open(data: bytes, *, values: bool) -> Any:
    return load_workbook(io.BytesIO(data), read_only=True, data_only=values)


def _sheet(book: Any, sheet: str | None, where: str) -> Any:
    """L'onglet voulu (le premier si ``None``) ; ni absent, ni graphique."""
    charts = [chart.title for chart in book.chartsheets]
    name = sheet if sheet is not None else next(iter(book.sheetnames), None)
    if name is None:
        raise ToolError(f"{where} : le classeur n'a aucun onglet de cellules.")
    if name in charts:
        raise ToolError(f"{where} : {name!r} est un onglet graphique, sans cellules.")
    if name not in book.sheetnames:
        known = ", ".join(book.sheetnames) or "aucun"
        raise ToolError(f"{where} : onglet {name!r} introuvable (onglets : {known}).")
    return book[name]


def _reset(ws: Any) -> None:
    """La dimension déclarée peut mentir : sans ce reset, une feuille annoncée « A1 »
    ne rendait que A1 (test « dimension »)."""
    ws.reset_dimensions()


def overview(data: bytes, where: str) -> dict[str, JsonValue]:
    """Les onglets : nom, état, dimension déclarée, et le début de chacun."""
    check_archive(data, where)
    deadline = time.monotonic() + READ_DEADLINE
    with _guarded(where):
        book = _open(data, values=True)
        try:
            sheets: list[JsonValue] = []
            for ws in book.worksheets:
                _in_time(deadline, where)
                declared = ws.calculate_dimension() if ws.max_row else None
                _reset(ws)
                sheets.append(
                    {
                        "name": ws.title,
                        "state": ws.sheet_state,
                        # Ce que le fichier annonce ; il peut mentir.
                        "dimension": declared,
                        "preview": cast(JsonValue, _preview(ws)),
                    }
                )
            charts = [sheet.title for sheet in book.chartsheets]
            return {"sheets": sheets, "chartsheets": cast(JsonValue, charts)}
        finally:
            book.close()


def _preview(ws: Any) -> list[list[JsonValue]]:
    rows = [
        [_json(value, PREVIEW_CHARS)[0] for value in values]
        for values in ws.iter_rows(
            min_row=1, max_row=PREVIEW_ROWS, max_col=PREVIEW_COLUMNS, values_only=True
        )
    ]
    return _trimmed(rows, [[value is not None for value in row] for row in rows])


def _bounds(cells: str | None) -> tuple[int, int, int, int]:
    """Première colonne, première ligne, dernière colonne, dernière ligne : dans l'ordre."""
    if cells is None:
        return 1, 1, MAX_SHEET_COLUMNS, MAX_SHEET_ROWS
    try:
        min_col, min_row, max_col, max_row = range_boundaries(cells.upper().replace("$", ""))
    except ValueError, TypeError:
        raise ToolError(f"{cells} : plage illisible (A1:F50, A:C, 2:10).") from None
    columns = [c for c in (min_col, max_col) if c is not None]
    rows = [r for r in (min_row, max_row) if r is not None]
    if any(not 1 <= r <= MAX_SHEET_ROWS for r in rows):
        raise ToolError(f"{cells} : une ligne va de 1 à {MAX_SHEET_ROWS}.")
    if any(not 1 <= c <= MAX_SHEET_COLUMNS for c in columns):
        raise ToolError(f"{cells} : une colonne va de A à XFD.")
    # Comme Excel : « C5:A1 » est la plage A1:C5.
    return (
        min(columns, default=1),
        min(rows, default=1),
        max(columns, default=MAX_SHEET_COLUMNS),
        max(rows, default=MAX_SHEET_ROWS),
    )


def read_range(
    data: bytes, where: str, sheet: str, cells: str | None, max_rows: int
) -> tuple[Read, str]:
    """Les cellules d'une plage ; sans plage, la feuille depuis A1. Et l'état de l'onglet."""
    check_archive(data, where)
    first_column, first_row, last_column, last_row = _bounds(cells)
    wide = last_column - first_column + 1 > MAX_COLUMNS
    shown_last = first_column + MAX_COLUMNS - 1 if wide else last_column
    wanted = min(last_row - first_row + 1, max_rows)
    deadline = time.monotonic() + READ_DEADLINE
    with _guarded(where):
        book = _open(data, values=False)
        try:
            ws = _sheet(book, sheet, where)
            state = ws.sheet_state
            declared_last = ws.max_column
            _reset(ws)
            # Une colonne de plus que la borne : savoir s'il y a des données à droite.
            reach = shown_last + 1 if wide else shown_last
            cells_rows = _iter(ws, 1, first_column, reach, last_row, deadline, where)
            raw: list[tuple[Any, ...]] = []
            used = 0
            budget_hit = False
            beyond = False
            for number, row in cells_rows:
                if number < first_row:
                    continue
                kept = row[: shown_last - first_column + 1]
                beyond = beyond or any(cell.value is not None for cell in row[len(kept) :])
                size = _size(kept)
                if raw and used + size > READ_CHARS:
                    # La ligne ne tient plus : elle sera la première de la suite.
                    budget_hit = True
                    break
                raw.append(kept)
                used += size
                if len(raw) == wanted:
                    break
            # Une ligne non vide plus bas, même après des lignes vides : la lecture est coupée.
            more = budget_hit or any(
                any(cell.value is not None for cell in row[: shown_last - first_column + 1])
                for _, row in cells_rows
            )
        finally:
            book.close()
        has_formula = any(_is_formula(cell) for row in raw for cell in row)
        cached = (
            _cached(data, sheet, first_row, first_column, shown_last, len(raw), deadline, where)
            if has_formula
            else [[cell.value for cell in row] for row in raw]
        )
    found = _read(raw, cached, first_row, first_column)
    columns_cut_after = (
        get_column_letter(shown_last)
        if wide and (beyond or (declared_last or 0) > shown_last)
        else None
    )
    return (
        Read(
            first_row=first_row,
            first_column=first_column,
            rows=found.rows,
            formulas=found.formulas,
            without_value=found.without_value,
            next_row=first_row + len(raw) if more else None,
            columns_cut_after=columns_cut_after,
            cut_cells=found.cut_cells,
            budget_hit=budget_hit,
        ),
        state,
    )


def _size(cells: tuple[Any, ...]) -> int:
    """La place d'une ligne dans le résultat, sans les cellules vides de droite."""
    values = [cell.value for cell in cells]
    while values and values[-1] is None:
        values.pop()
    return len(json.dumps([_json(value)[0] for value in values], default=str))


def _cached(
    data: bytes,
    sheet: str,
    first_row: int,
    first_column: int,
    last_column: int,
    count: int,
    deadline: float,
    where: str,
) -> list[list[Any]]:
    """Les valeurs en cache du classeur : une seconde passe, seulement s'il a des formules."""
    _in_time(deadline, where)
    book = _open(data, values=True)
    try:
        ws = _sheet(book, sheet, where)
        _reset(ws)
        last_row = first_row + max(count, 1) - 1
        return [
            [cell.value for cell in row]
            for number, row in _iter(ws, 1, first_column, last_column, last_row, deadline, where)
            if number >= first_row
        ][:count]
    finally:
        book.close()


def _iter(
    ws: Any,
    first_row: int,
    first_column: int,
    last_column: int,
    last_row: int,
    deadline: float,
    where: str,
) -> Iterator[tuple[int, tuple[Any, ...]]]:
    """Les lignes numérotées depuis ``first_row``, l'échéance vérifiée toutes les 1 000 lignes.

    Depuis la ligne 1 pour une plage plus bas : openpyxl parcourt les lignes
    avant ``min_row`` sans rendre la main ; compter soi-même permet de s'arrêter.
    """
    rows = ws.iter_rows(
        min_row=first_row, max_row=last_row, min_col=first_column, max_col=last_column
    )
    for number, row in enumerate(rows, start=first_row):
        if number % 1_000 == 0:
            _in_time(deadline, where)
        yield number, row


def _in_time(deadline: float, where: str) -> None:
    """Un fil d'exécution ne s'annule pas : la lecture s'arrête d'elle-même à l'échéance."""
    if time.monotonic() > deadline:
        raise ToolError(
            f"{where} : lecture arrêtée après {READ_DEADLINE:g} s ; viser une plage plus "
            "proche du début, ou un classeur plus petit."
        )


@dataclass(frozen=True, slots=True)
class _Found:
    rows: list[list[JsonValue]]
    formulas: dict[str, str]
    without_value: int
    cut_cells: int


def _read(
    raw: list[tuple[Any, ...]], cached: list[list[Any]], first_row: int, first_column: int
) -> _Found:
    rows: list[list[JsonValue]] = []
    written: list[list[bool]] = []
    formulas: dict[str, str] = {}
    without_value = 0
    cut = 0
    for index, (cells, values) in enumerate(zip(raw, cached, strict=False)):
        row: list[JsonValue] = []
        mask: list[bool] = []
        for offset, (cell, value) in enumerate(zip(cells, values, strict=False)):
            if _is_formula(cell):
                address = f"{get_column_letter(first_column + offset)}{first_row + index}"
                formulas[address] = _formula_text(cell.value)
                without_value += value is None
                shown, was_cut = _json(value)
            else:
                shown, was_cut = _json(cell.value)
            cut += was_cut
            row.append(shown)
            # Une formule sans valeur compte comme écrite : elle ne disparaît pas des lignes.
            mask.append(cell.value is not None)
        rows.append(row)
        written.append(mask)
    return _Found(_trimmed(rows, written), formulas, without_value, cut)


def _is_formula(cell: Any) -> bool:
    """Une formule d'après le type de la cellule, pas d'après un « = » en tête de texte."""
    return getattr(cell, "data_type", None) == "f"


def _formula_text(value: object) -> str:
    if isinstance(value, ArrayFormula):
        # Une formule matricielle : son texte et la plage qu'elle remplit.
        return f"{{{value.text}}} sur {value.ref}"
    return str(value)


def _json(value: object, limit: int = MAX_CELL_CHARS) -> tuple[JsonValue, bool]:
    """Une valeur de cellule en JSON ; vrai si un texte a été coupé."""
    match value:
        case None | bool() | int():
            return value, False
        case float():
            if not math.isfinite(value):
                # JSON n'a ni infini ni NaN : le texte, plutôt qu'un null muet.
                return str(value), False
            return (int(value) if value.is_integer() and abs(value) < 2**53 else value), False
        case _:
            text = _text(value)
            if len(text) > limit:
                return text[:limit] + "…", True
            return text, False


def _text(value: object) -> str:
    match value:
        case dt.datetime():
            if value.time() == dt.time():
                return value.date().isoformat()
            return value.isoformat(timespec="seconds")
        case dt.date() | dt.time():
            return value.isoformat()
        case dt.timedelta():
            hours, rest = divmod(int(value.total_seconds()), 3600)
            return f"{hours}:{rest // 60:02}:{rest % 60:02}"
        case _:
            return str(value)


def _trimmed(rows: list[list[JsonValue]], written: list[list[bool]]) -> list[list[JsonValue]]:
    """Sans les lignes vides de la fin, ni les colonnes vides de droite (selon ce qui est écrit)."""
    while rows and not any(written[len(rows) - 1]):
        rows.pop()
    width = max(
        (max((i + 1 for i, w in enumerate(mask) if w), default=0) for mask in written[: len(rows)]),
        default=0,
    )
    return [row[:width] for row in rows]


def records(data: bytes, where: str, sheet: str | None) -> list[dict[str, JsonValue]]:
    """Un onglet comme source de ``data`` : la ligne 1 porte les en-têtes.

    Les valeurs passent entières (pas de coupe à 1 000 caractères) ; une formule
    sans valeur en cache est refusée plutôt que lue comme une cellule vide.
    """
    check_archive(data, where)
    deadline = time.monotonic() + READ_DEADLINE
    with _guarded(where):
        book = _open(data, values=False)
        try:
            ws = _sheet(book, sheet, where)
            _reset(ws)
            name = ws.title
            raw = _source_rows(ws, deadline, where)
        finally:
            book.close()
        values = raw
        if any(_is_formula(cell) for row in raw for cell in row):
            cached_book = _open(data, values=True)
            try:
                cached_ws = _sheet(cached_book, name, where)
                _reset(cached_ws)
                values = _source_rows(cached_ws, deadline, where)
            finally:
                cached_book.close()
    headers = [
        None if cell.value is None else _text(cell.value) for cell in (raw[0] if raw else ())
    ]
    named = [h for h in headers if h is not None]
    doubles = sorted({h for h in named if named.count(h) > 1})
    if doubles:
        raise ToolError(f"{where} : en-têtes en double : {', '.join(doubles)}.")
    found: list[dict[str, JsonValue]] = []
    missing = 0
    for cells, cached in zip(raw[1:], values[1:], strict=False):
        if all(cell.value is None for cell in cells):
            continue
        if len(found) == MAX_DATA_ROWS:
            raise ToolError(f"{where} : plus de {MAX_DATA_ROWS} lignes, au-delà de la borne.")
        record: dict[str, JsonValue] = {}
        for header, cell, value in zip(headers, cells, cached, strict=False):
            if header is None:
                continue
            if _is_formula(cell) and value.value is None:
                missing += 1
            record[header] = _plain(value.value)
        found.append(record)
    if missing:
        raise ToolError(
            f"{where} : {missing} formule(s) sans valeur calculée dans la source ; l'ouvrir et "
            "l'enregistrer dans Excel ou LibreOffice, ou la donner en CSV."
        )
    return found


def _source_rows(ws: Any, deadline: float, where: str) -> list[tuple[Any, ...]]:
    """En-tête et lignes, bornés : au-delà de la borne d'une source, inutile de tout lire."""
    rows: list[tuple[Any, ...]] = []
    for _, row in _iter(ws, 1, 1, MAX_COLUMNS, MAX_SHEET_ROWS, deadline, where):
        rows.append(row)
        if len(rows) > MAX_DATA_ROWS + 1:
            raise ToolError(f"{where} : plus de {MAX_DATA_ROWS} lignes, au-delà de la borne.")
    return rows


def _plain(value: object) -> JsonValue:
    """Une valeur de source, entière : nombres et booléens tels quels, le reste en texte."""
    match value:
        case None | bool() | int():
            return value
        case float():
            return value if math.isfinite(value) else str(value)
        case _:
            return _text(value)
