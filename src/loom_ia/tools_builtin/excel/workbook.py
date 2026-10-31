# SPDX-License-Identifier: Apache-2.0
"""Le classeur : description validée, puis mise en forme par le style maison, puis octets.

Le style est figé ici : police, couleurs par rôle, formats par type, largeurs,
volets gelés, filtre, légende ; les graphiques, dans ``charts``. Deux descriptions
identiques donnent le même fichier, octet pour octet : les dates du classeur et
de l'archive sont fixées.
"""

import datetime as dt
import io
import zipfile
from collections.abc import Mapping
from dataclasses import dataclass, field
from typing import Final
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.writer.excel import ExcelWriter

from loom_ia.tools_builtin._values import Stored, convert
from loom_ia.tools_builtin.excel import charts
from loom_ia.tools_builtin.excel.arguments import (
    MAX_COLUMNS,
    MAX_DATA_ROWS,
    MAX_INLINE_ROWS,
    CellSpec,
    ColumnType,
    ExcelArguments,
    Problems,
    RowSpec,
    SheetSpec,
)
from loom_ia.tools_builtin.excel.formulas import AGGREGATES, check_formula, sheet_name_problem

# Dates du classeur (docProps/core.xml) et de l'archive : fixes, pour que deux
# appels identiques donnent le même fichier (reprise idempotente, empreinte).
FIXED_DATE: Final = dt.datetime(2000, 1, 1)
ZIP_DATE: Final = (2000, 1, 1, 0, 0, 0)

# --- Style maison --------------------------------------------------------------------

FONT_NAME: Final = "Calibri"
FONT_SIZE: Final = 11
BLUE: Final = "0000FF"
GREEN: Final = "008000"
GREY: Final = "7F7F7F"
HEADER_FILL: Final = "D9E1F2"
INPUT_FILL: Final = "FFF2CC"
ASSUMPTION_FILL: Final = "FFFF99"
MIN_WIDTH: Final = 8
MAX_WIDTH: Final = 60
# Largeur totale (en caractères) au-delà de laquelle un onglet s'imprime en paysage :
# à peu près ce que tient une page A4 en portrait, marges comprises.
PORTRAIT_WIDTH: Final = 90
# Types centrés : une date alignée à droite colle au texte de la colonne suivante.
CENTERED: Final[frozenset[ColumnType]] = frozenset({"date", "date_heure", "annee", "booleen"})

# Formats en syntaxe neutre : Excel applique les séparateurs de la locale du
# lecteur. Dates : formats intégrés n° 14 et 22, qui suivent la locale (D5 :
# adaptatif) ; négatifs avec le signe moins, zéro affiché 0.
NUMBER_FORMATS: Final[Mapping[ColumnType, str]] = {
    "texte": "@",
    "entier": "#,##0",
    "decimal": "#,##0.00",
    "montant_eur": '#,##0.00 "€"',
    "montant_keur": '#,##0 "k€"',
    "pourcentage": "0.0%",
    "date": "mm-dd-yy",
    "date_heure": "m/d/yy h:mm",
    "annee": "@",
    "booleen": "@",
}
# Largeur d'une colonne de formules, dont la valeur n'est connue qu'au recalcul.
FORMULA_WIDTH: Final[Mapping[ColumnType, int]] = {
    "texte": 20,
    "entier": 10,
    "decimal": 12,
    "montant_eur": 14,
    "montant_keur": 12,
    "pourcentage": 9,
    "date": 11,
    "date_heure": 16,
    "annee": 6,
    "booleen": 6,
}

_THIN: Final = Side(style="thin", color="A6A6A6")
_HEADER_STYLE: Final = {
    "font": Font(name=FONT_NAME, size=FONT_SIZE, bold=True),
    "fill": PatternFill("solid", fgColor=HEADER_FILL),
    "border": Border(bottom=_THIN),
    "alignment": Alignment(horizontal="center", vertical="center", wrap_text=True),
}


@dataclass(slots=True)
class Written:
    """Ce que l'on retient d'un onglet écrit, pour le résumé et le recalcul."""

    name: str
    rows: int
    formulas: list[str] = field(default_factory=list[str])
    charts: int = 0


@dataclass(frozen=True, slots=True)
class Prepared:
    """La description validée, ses lignes chargées : prête à écrire."""

    sheets: tuple[tuple[SheetSpec, list[RowSpec]], ...]
    warnings: tuple[str, ...]
    # Le fuseau où passent les dates-heures qui en portent un.
    zone: ZoneInfo
    # Les graphiques validés de chaque onglet, dans l'ordre des onglets.
    charts: tuple[tuple[charts.Plotted, ...], ...] = ()


def prepare(
    arguments: ExcelArguments, sources: Mapping[str, list[RowSpec]], zone: ZoneInfo
) -> Prepared:
    """Valide la description entière ; ``ToolError`` qui liste toutes les fautes trouvées.

    ``sources`` : les lignes lues des fichiers, par chemin donné.
    """
    problems = Problems()
    names: dict[str, str] = {}
    for sheet in arguments.sheets:
        problem = sheet_name_problem(sheet.name)
        if problem:
            problems.error(problem)
        elif sheet.name.lower() in names:
            problems.error(f"onglet {sheet.name!r} donné deux fois (casse ignorée)")
        else:
            names[sheet.name.lower()] = sheet.name
    prepared: list[tuple[SheetSpec, list[RowSpec]]] = []
    plotted: list[tuple[charts.Plotted, ...]] = []
    for sheet in arguments.sheets:
        rows = _sheet_rows(sheet, sources, problems)
        _check_sheet(sheet, rows, problems)
        prepared.append((sheet, rows))
        plotted.append(tuple(charts.check(sheet, rows, problems)))
    problems.raise_if_any()
    return Prepared(
        sheets=tuple(prepared),
        warnings=tuple(problems.warnings),
        zone=zone,
        charts=tuple(plotted),
    )


def _sheet_rows(
    sheet: SheetSpec, sources: Mapping[str, list[RowSpec]], problems: Problems
) -> list[RowSpec]:
    if sheet.data is not None:
        return sources.get(sheet.data.path, [])
    rows = sheet.rows or []
    if len(rows) > MAX_INLINE_ROWS:
        problems.error(
            f"{sheet.name} : {len(rows)} lignes données en ligne, au-delà de {MAX_INLINE_ROWS} "
            "— les écrire dans un fichier CSV ou JSON et le citer dans 'data'"
        )
    return rows


def _check_sheet(sheet: SheetSpec, rows: list[RowSpec], problems: Problems) -> None:
    headers = [column.header for column in sheet.columns]
    if len(headers) > MAX_COLUMNS:
        problems.error(f"{sheet.name} : {len(headers)} colonnes, au-delà de {MAX_COLUMNS}")
    doubles = sorted({h for h in headers if headers.count(h) > 1})
    if doubles:
        problems.error(f"{sheet.name} : en-tête donné deux fois : {', '.join(doubles)}")
    if len(rows) > MAX_DATA_ROWS:
        problems.error(f"{sheet.name} : {len(rows)} lignes, au-delà de {MAX_DATA_ROWS}")
    for number, row in enumerate(rows, start=2):
        if len(row.cells) > len(headers):
            problems.error(
                f"{sheet.name}!{number} : {len(row.cells)} valeurs pour {len(headers)} colonnes"
            )
        if (
            row.role == "hypothese"
            and not row.note
            and not any(isinstance(cell, CellSpec) and cell.note for cell in row.cells)
        ):
            problems.error(
                f"{sheet.name}!{number} : une hypothèse porte sa source en 'note' (« Source : … »)"
            )
    for header in sheet.totals or {}:
        if header not in headers:
            problems.error(f"{sheet.name} : total sur une colonne inconnue {header!r}")
    if sheet.totals and not rows:
        problems.warn(f"{sheet.name} : aucune ligne, la ligne de total n'est pas écrite")


def build_workbook(prepared: Prepared) -> tuple[Workbook, list[Written], list[str]]:
    """Le classeur mis en forme ; les formules sont déjà validées par ``prepare``."""
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    names = {sheet.name.lower(): sheet.name for sheet, _ in prepared.sheets}
    problems = Problems(warnings=list(prepared.warnings))
    written: list[Written] = []
    for number, (sheet, rows) in enumerate(prepared.sheets):
        ws = workbook.create_sheet(sheet.name)
        done = _write_sheet(ws, sheet, rows, names, problems, prepared.zone)
        plotted = prepared.charts[number] if prepared.charts else ()
        if plotted:
            charts.draw(ws, sheet, plotted)
            done.charts = len(plotted)
        written.append(done)
    problems.raise_if_any()
    workbook.properties.creator = "loom-ia"
    workbook.properties.lastModifiedBy = "loom-ia"
    # Formules sans valeur en cache : Excel recalcule tout à l'ouverture.
    workbook.calculation.fullCalcOnLoad = True
    return workbook, written, problems.warnings


def _write_sheet(
    ws: Worksheet,
    sheet: SheetSpec,
    rows: list[RowSpec],
    names: Mapping[str, str],
    problems: Problems,
    zone: ZoneInfo,
) -> Written:
    done = Written(name=sheet.name, rows=len(rows))
    columns = sheet.columns
    widths = [len(column.header) for column in columns]
    for index, column in enumerate(columns, start=1):
        cell = _cell(ws, 1, index, column.header)
        for attribute, style in _HEADER_STYLE.items():
            setattr(cell, attribute, style)
    if any(len(column.header) > MAX_WIDTH for column in columns):
        ws.row_dimensions[1].height = 30
    has_input = has_assumption = False
    for number, row in enumerate(rows, start=2):
        noted = False
        for index, given in enumerate(row.cells, start=1):
            column = columns[index - 1]
            spec = given if isinstance(given, CellSpec) else CellSpec(value=given)
            where = f"{sheet.name}!{get_column_letter(index)}{number}"
            cell = _cell(ws, number, index)
            role = spec.role or row.role
            if isinstance(spec.value, str) and spec.value.startswith("="):
                cell.value = check_formula(spec.value, where, names, problems)
                done.formulas.append(where)
                role = "total" if role == "total" else None
                _style(cell, role or ("lien" if "!" in spec.value else "formule"), column.type)
                widths[index - 1] = max(widths[index - 1], FORMULA_WIDTH[column.type])
            else:
                cell.value = convert(spec.value, column.type, where, problems, zone)
                # Un texte non vide d'une ligne à saisir est son libellé, pas une saisie.
                label = column.type == "texte" and cell.value is not None
                if role in {"saisie", "hypothese"} and label and spec.role is None:
                    role = None
                _style(cell, role, column.type)
                widths[index - 1] = max(widths[index - 1], len(_shown(cell.value)))
            has_input |= role == "saisie"
            has_assumption |= role == "hypothese"
            note = spec.note or (row.note if not noted and role not in {None, "note"} else None)
            if note:
                cell.comment = Comment(note, "loom-ia", width=300, height=80)
                noted = noted or note == row.note
        if row.note and not noted and row.cells:
            _cell(ws, number, 1).comment = Comment(row.note, "loom-ia", width=300, height=80)
        if row.role in {"total", "note"}:
            for index in range(len(row.cells) + 1, len(columns) + 1):
                _style(_cell(ws, number, index), row.role, columns[index - 1].type)
    last = len(rows) + 1
    if sheet.totals and rows:
        last += 1
        _write_totals(ws, sheet, last, done)
    _finish(ws, sheet, len(rows), last, widths, has_input, has_assumption)
    return done


def _write_totals(ws: Worksheet, sheet: SheetSpec, at: int, done: Written) -> None:
    totals = sheet.totals or {}
    label = _total_label(sheet)
    for index, column in enumerate(sheet.columns, start=1):
        cell = _cell(ws, at, index)
        aggregate = totals.get(column.header)
        if aggregate is not None:
            letter = get_column_letter(index)
            cell.value = f"={AGGREGATES[aggregate]}({letter}2:{letter}{at - 1})"
            done.formulas.append(f"{sheet.name}!{letter}{at}")
        elif index == label:
            cell.value = "Total"
            _style(cell, "total", "texte")
            continue
        _style(cell, "total", column.type)


def _total_label(sheet: SheetSpec) -> int | None:
    """La colonne (à partir de 1) où écrire « Total » : la première sans agrégat, texte d'abord.

    Une date en première colonne ne laisse plus le total sans libellé ; ``None``
    si toutes les colonnes ont leur agrégat.
    """
    totals = sheet.totals or {}
    free = [
        (index, column)
        for index, column in enumerate(sheet.columns, start=1)
        if column.header not in totals
    ]
    texts = [index for index, column in free if column.type == "texte"]
    if texts:
        return texts[0]
    return free[0][0] if free else None


def _finish(
    ws: Worksheet,
    sheet: SheetSpec,
    rows: int,
    last: int,
    widths: list[int],
    has_input: bool,
    has_assumption: bool,
) -> None:
    ws.freeze_panes = "A2"
    end = get_column_letter(len(sheet.columns))
    if rows:
        ws.auto_filter.ref = f"A1:{end}{rows + 1}"
    for index, (width, column) in enumerate(zip(widths, sheet.columns, strict=True), start=1):
        # Une marge, et le retrait d'un texte (d'un caractère) en plus.
        margin = 3 if column.type == "texte" else 2
        ws.column_dimensions[get_column_letter(index)].width = min(
            MAX_WIDTH, max(MIN_WIDTH, width + margin)
        )
    legend: list[str] = []
    if has_input:
        legend.append("Bleu sur fond jaune pâle : cellule à saisir.")
    if has_assumption:
        legend.append("Bleu sur fond jaune : hypothèse (sa source en commentaire).")
    for offset, line in enumerate(legend):
        cell = _cell(ws, last + 2 + offset, 1, line)
        _style(cell, "note", "texte")
    _page_setup(ws)


def _page_setup(ws: Worksheet) -> None:
    """À l'impression et à l'export PDF : une page de large, en paysage si l'onglet est large.

    Sans ce réglage, une colonne de trop part seule sur une deuxième page.
    """
    total = sum(dimension.width or 0 for dimension in ws.column_dimensions.values())
    ws.page_setup.orientation = "landscape" if total > PORTRAIT_WIDTH else "portrait"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    # 0 : autant de pages en hauteur qu'il en faut.
    ws.page_setup.fitToHeight = 0


def _cell(ws: Worksheet, row: int, column: int, value: Stored = None) -> Cell:
    """Une cellule à écrire : jamais fusionnée, l'outil ne fusionne rien."""
    found = ws.cell(row=row, column=column)
    assert isinstance(found, Cell)
    if value is not None:
        found.value = value
    return found


def _style(cell: Cell, role: str | None, kind: ColumnType) -> None:
    cell.number_format = NUMBER_FORMATS[kind]
    match role:
        case "saisie":
            cell.font = Font(name=FONT_NAME, size=FONT_SIZE, color=BLUE)
            cell.fill = PatternFill("solid", fgColor=INPUT_FILL)
        case "hypothese":
            cell.font = Font(name=FONT_NAME, size=FONT_SIZE, color=BLUE)
            cell.fill = PatternFill("solid", fgColor=ASSUMPTION_FILL)
        case "lien":
            cell.font = Font(name=FONT_NAME, size=FONT_SIZE, color=GREEN)
        case "total":
            cell.font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True)
            cell.border = Border(top=_THIN)
        case "note":
            cell.font = Font(name=FONT_NAME, size=FONT_SIZE, italic=True, color=GREY)
        case _:
            cell.font = Font(name=FONT_NAME, size=FONT_SIZE)
    if role == "note":
        return
    if kind in CENTERED:
        cell.alignment = Alignment(horizontal="center")
    elif kind == "texte":
        # Un retrait : un texte ne colle pas au nombre aligné à droite de la colonne voisine.
        wrap = isinstance(cell.value, str) and len(cell.value) > MAX_WIDTH
        cell.alignment = Alignment(
            horizontal="left", indent=1, wrap_text=wrap, vertical="top" if wrap else None
        )


def _shown(value: object) -> str:
    match value:
        case None:
            return ""
        case float():
            return f"{value:,.2f}"
        case int():
            return f"{value:,}"
        case dt.datetime():
            return "00/00/0000 00:00"
        case dt.date():
            return "00/00/0000"
        case _:
            return str(value)


def to_bytes(workbook: Workbook) -> bytes:
    """Le fichier .xlsx, le même octet pour octet pour un même classeur."""
    workbook.properties.created = FIXED_DATE
    workbook.properties.modified = FIXED_DATE
    raw = io.BytesIO()
    # ``ExcelWriter`` et non ``Workbook.save`` : ce dernier date le fichier de maintenant.
    with zipfile.ZipFile(raw, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as archive:
        ExcelWriter(workbook, archive).save()
    fixed = io.BytesIO()
    with (
        zipfile.ZipFile(raw) as source,
        zipfile.ZipFile(fixed, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for info in source.infolist():
            entry = zipfile.ZipInfo(info.filename, date_time=ZIP_DATE)
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = 0o600 << 16
            target.writestr(entry, source.read(info.filename))
    return fixed.getvalue()
