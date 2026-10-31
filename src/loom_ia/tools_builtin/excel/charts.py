# SPDX-License-Identifier: Apache-2.0
"""Les graphiques d'un onglet : validés avec la description du classeur, puis dessinés.

Un graphique cite les colonnes de son onglet par leur en-tête : ses libellés
(``categories``) et ses séries (``values``), des colonnes de nombres d'un même
type. Il trace les lignes de données, depuis la ligne 2 : ni la ligne de
``totals``, ni les lignes ``total`` ou ``note`` de la fin. Les fautes sont dites
avec celles du reste du classeur.

Les bornes sont celles des documents (``word``, ``pdf``), sauf pour une courbe,
qui se lit encore avec beaucoup de points.

Un graphique est celui d'Excel, **lié aux cellules** de son onglet : ses
libellés et ses séries sont des plages (``'Ventes'!$A$2:$A$7``), pas des
nombres recopiés. Changer une cellule change le graphique, et « Sélectionner
les données » montre ces plages. Une colonne de formules se trace comme une
autre : Excel la calcule à l'ouverture.

Le style est celui des graphiques de ``word`` et ``pdf`` (``_chart_style``),
appliqué par l'API d'openpyxl. Le graphique se place à droite des données,
depuis la ligne 2 ; plusieurs graphiques s'empilent.
"""

import math
from collections.abc import Sequence
from dataclasses import dataclass
from typing import Final, cast

from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout
from openpyxl.chart.legend import Legend
from openpyxl.chart.marker import DataPoint, Marker
from openpyxl.chart.series import Series
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Font, Paragraph, ParagraphProperties
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from loom_ia.tools_builtin._chart_style import (
    AXIS_FORMATS,
    GRID,
    LABEL_FORMATS,
    LABEL_ROOM,
    LABELLED_BARS,
    PALETTE,
    TEXT,
    axis_end,
)
from loom_ia.tools_builtin._document.arguments import MAX_CATEGORIES, MAX_SERIES, MAX_SLICES
from loom_ia.tools_builtin._document.content import NUMERIC
from loom_ia.tools_builtin._values import ColumnType
from loom_ia.tools_builtin.excel.arguments import (
    CellSpec,
    ChartSpec,
    Problems,
    RowSpec,
    SheetSpec,
)

# Une courbe se lit encore avec beaucoup de points (une année au jour le jour) ;
# des barres ou des secteurs, non.
MAX_POINTS: Final = 1_000
MAX_CHARTS: Final = 10
KIND_SAID: Final[dict[str, str]] = {
    "barres": "des barres",
    "barres_horizontales": "des barres",
    "courbe": "une courbe",
    "secteurs": "des secteurs",
}
# Des libellés que l'axe garde en nombres (et dont il garde le format) ; les autres sont du texte.
NUMBERED: Final[frozenset[ColumnType]] = NUMERIC | frozenset[ColumnType]({"date", "date_heure"})
# Taille d'un graphique, en centimètres ; une ligne d'Excel fait 0,53 cm (15 points).
WIDTH: Final = 16
HEIGHT: Final = 9
ROW_CM: Final = 0.53
# Les graphiques que l'outil dessine ; ceux qui ont des axes.
type Drawn = BarChart | LineChart | PieChart
type Axed = BarChart | LineChart
# Des lignes qui ne sont pas des données : elles ne font pas partie du graphique.
NOT_DATA: Final = frozenset({"total", "note"})


@dataclass(frozen=True, slots=True)
class Plotted:
    """Un graphique validé : ses colonnes (à partir de 1) et ce que l'on sait de ses valeurs."""

    spec: ChartSpec
    category: int
    values: tuple[int, ...]
    value_type: ColumnType
    category_type: ColumnType
    # Les lignes de données tracées : de la ligne 2 à ``last``.
    last: int
    # Les valeurs écrites en clair (sans formule) : bornes de l'axe, contrôle des secteurs.
    known: tuple[float, ...]
    # Toutes les valeurs sont connues (aucune formule) ?
    complete: bool


def data_rows(rows: Sequence[RowSpec]) -> int:
    """Combien de lignes, depuis la ligne 2, sont des données : avant tout total ou note."""
    for index, row in enumerate(rows):
        if row.role in NOT_DATA:
            return index
    return len(rows)


def check(sheet: SheetSpec, rows: Sequence[RowSpec], problems: Problems) -> list[Plotted]:
    """Les graphiques de l'onglet, validés ; chaque faute est dite à ``problems``."""
    specs = sheet.chart_specs()
    if not specs:
        return []
    if len(specs) > MAX_CHARTS:
        problems.error(f"{sheet.name} : {len(specs)} graphiques, au-delà de {MAX_CHARTS}")
    headers = [column.header.strip() for column in sheet.columns]
    count = data_rows(rows)
    after = [n + 2 for n, row in enumerate(rows[count:], start=count) if row.role not in NOT_DATA]
    plotted: list[Plotted] = []
    for number, spec in enumerate(specs, start=1):
        where = f"{sheet.name}, graphique {number}"
        if after:
            problems.error(
                f"{where} : la ligne {after[0]} suit une ligne de total ou de note — les données "
                "d'un graphique se suivent depuis la ligne 2, totaux et notes à la fin"
            )
            continue
        found = _check_one(spec, sheet, headers, rows[:count], where, problems)
        if found is not None:
            plotted.append(found)
    return plotted


def _check_one(
    spec: ChartSpec,
    sheet: SheetSpec,
    headers: list[str],
    rows: Sequence[RowSpec],
    where: str,
    problems: Problems,
) -> Plotted | None:
    wanted = [spec.categories, *spec.values]
    unknown = [name for name in wanted if name.strip() not in headers]
    if unknown:
        problems.error(
            f"{where} : colonne(s) inconnue(s) : {', '.join(unknown)} "
            f"(colonnes : {', '.join(headers)})"
        )
        return None
    category = headers.index(spec.categories.strip())
    values = [headers.index(name.strip()) for name in spec.values]
    if category in values or len(set(values)) != len(values):
        problems.error(f"{where} : une colonne citée deux fois")
        return None
    kinds = {sheet.columns[index].type for index in values}
    kind = next(iter(sorted(kinds)))
    if len(kinds) > 1 or kind not in NUMERIC:
        problems.error(
            f"{where} : les séries viennent de colonnes de nombres d'un même type "
            f"(ici : {', '.join(sorted(kinds))})"
        )
        return None
    if len(values) > MAX_SERIES:
        problems.error(f"{where} : {len(values)} séries, au-delà de {MAX_SERIES}")
    count = len(rows)
    limit = MAX_POINTS if spec.kind == "courbe" else MAX_CATEGORIES
    if not count:
        problems.error(f"{where} : aucune ligne de données, le graphique serait vide")
    elif count > limit:
        problems.error(
            f"{where} : {count} lignes, au-delà de {limit} pour {KIND_SAID[spec.kind]} — "
            "regrouper, ou garder les principales"
        )
    known: list[float] = []
    complete = True
    for row in rows:
        for index in values:
            number = _literal(row.cells[index] if index < len(row.cells) else None)
            if number is None:
                complete = False
            else:
                known.append(number)
    if spec.kind == "secteurs":
        if len(values) > 1:
            problems.error(f"{where} : un graphique en secteurs n'a qu'une série")
        if count > MAX_SLICES:
            problems.error(
                f"{where} : {count} secteurs, au-delà de {MAX_SLICES} — illisible ; garder les "
                "principaux et regrouper le reste en « Autres », ou des barres"
            )
        if any(value < 0 for value in known):
            problems.error(f"{where} : une valeur négative n'a pas de secteur")
        elif complete and known and not any(known):
            problems.error(f"{where} : toutes les valeurs sont nulles")
    return Plotted(
        spec=spec,
        category=category + 1,
        values=tuple(index + 1 for index in values),
        value_type=kind,
        category_type=sheet.columns[category].type,
        last=count + 1,
        known=tuple(known),
        complete=complete,
    )


def _literal(given: object) -> float | None:
    """La valeur d'une cellule écrite en clair ; ``None`` pour une formule, un texte, un vide."""
    value = given.value if isinstance(given, CellSpec) else given
    if isinstance(value, bool) or not isinstance(value, int | float):
        return None
    return float(value)


# --- Dessin ------------------------------------------------------------------------------


def draw(ws: Worksheet, sheet: SheetSpec, plotted: Sequence[Plotted]) -> None:
    """Les graphiques, à droite des données, empilés depuis la ligne 2."""
    column = get_column_letter(len(sheet.columns) + 2)
    row = 2
    for found in plotted:
        chart = _chart(ws, found)
        ws.add_chart(chart, f"{column}{row}")
        row += math.ceil(float(chart.height) / ROW_CM) + 1


def _chart(ws: Worksheet, found: Plotted) -> Drawn:
    spec = found.spec
    labels = Reference(ws, min_col=found.category, min_row=2, max_row=found.last)
    chart: Drawn
    match spec.kind:
        case "secteurs":
            chart = PieChart()
            chart.varyColors = True
        case "courbe":
            chart = LineChart()
            chart.grouping = "standard"
        case "barres" | "barres_horizontales":
            bars = BarChart()
            bars.type = "bar" if spec.kind == "barres_horizontales" else "col"
            bars.grouping = "clustered"
            bars.gapWidth = 60
            chart = bars
    for index in found.values:
        # La ligne 1 (l'en-tête) donne son nom à la série.
        data = Reference(ws, min_col=index, min_row=1, max_row=found.last)
        chart.add_data(data, titles_from_data=True)  # pyright: ignore[reportUnknownMemberType]
    chart.set_categories(labels)  # pyright: ignore[reportUnknownMemberType]
    if found.category_type not in NUMBERED:
        # Des libellés en texte : une plage de textes, pas de nombres (``set_categories``).
        for series in _series(chart):
            series.cat = AxDataSource(strRef=StrRef(f=labels.__str__()))
    chart.width = WIDTH
    chart.height = _height(spec, found.last - 1)
    chart.style = None
    chart.roundedCorners = False
    chart.display_blanks = "gap"
    chart.title = spec.title
    if spec.title:
        _title_style(chart)
    several = len(found.values) > 1
    chart.legend = (
        Legend(legendPos="r" if spec.kind == "secteurs" else "b", txPr=_text(900))
        if spec.kind == "secteurs" or several
        else None
    )
    label = LABEL_FORMATS[found.value_type]
    match chart:
        case PieChart():
            _pie(chart, found.last - 1)
        case LineChart():
            _lines(chart)
            _axes(chart, found)
        case BarChart():
            _bars(chart, found, label)
            _axes(chart, found)
    return chart


def _series(chart: Drawn) -> list[Series]:
    # Le stub d'openpyxl ne type pas ``series`` (un alias de ``ser``).
    return list(cast(Sequence[Series], chart.series))  # pyright: ignore[reportUnknownMemberType]


def _height(spec: ChartSpec, count: int) -> int:
    """Des barres horizontales prennent la place de leurs libellés, comme dans Word."""
    if spec.kind != "barres_horizontales":
        return HEIGHT
    return min(20, max(6, round(1.4 + 0.6 * count)))


def _text(size: int, *, bold: bool = False, color: str = TEXT) -> RichText:
    props = CharacterProperties(sz=size, b=bold, solidFill=color, latin=Font(typeface="Calibri"))
    # ``r=[]`` : sans lui, openpyxl ajoute un passage de texte vide.
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=props), endParaRPr=props, r=[])])


def _title_style(chart: Drawn) -> None:
    title = chart.title
    assert title is not None and not isinstance(title, str)
    rich = title.tx.rich if title.tx is not None else None
    if rich is not None:
        for paragraph in rich.p:
            props = CharacterProperties(
                sz=1100, b=True, solidFill=PALETTE[0], latin=Font(typeface="Calibri")
            )
            paragraph.pPr = ParagraphProperties(defRPr=props)
            for run in paragraph.r:
                run.rPr = props
    title.overlay = False


def _labelled(found: Plotted) -> bool:
    return len(found.values) == 1 and found.last - 1 <= LABELLED_BARS


def _bars(chart: BarChart, found: Plotted, label: str) -> None:
    for index, series in enumerate(_series(chart)):
        color = PALETTE[index % len(PALETTE)]
        series.graphicalProperties = GraphicalProperties(solidFill=color)
        series.invertIfNegative = False
    if _labelled(found):
        chart.dataLabels = DataLabelList(
            numFmt=label, showVal=True, showLegendKey=False, showCatName=False,
            showSerName=False, showPercent=False, showBubbleSize=False,
            dLblPos="outEnd", txPr=_text(800),
        )  # fmt: skip


def _lines(chart: LineChart) -> None:
    for index, series in enumerate(_series(chart)):
        color = PALETTE[index % len(PALETTE)]
        series.graphicalProperties = GraphicalProperties(
            ln=LineProperties(solidFill=color, w=28575, cap="rnd")
        )
        series.marker = Marker(symbol="circle", size=5, spPr=GraphicalProperties(solidFill=color))
        series.smooth = False


def _pie(chart: PieChart, count: int) -> None:
    series = _series(chart)[0]
    series.dPt = [
        DataPoint(
            idx=index,
            spPr=GraphicalProperties(
                solidFill=PALETTE[index % len(PALETTE)],
                ln=LineProperties(solidFill="FFFFFF", w=12700),
            ),
        )
        for index in range(count)
    ]
    chart.dataLabels = DataLabelList(
        numFmt="0%", showVal=False, showPercent=True, showLegendKey=False,
        showCatName=False, showSerName=False, showBubbleSize=False,
        dLblPos="outEnd", txPr=_text(800),
    )  # fmt: skip
    chart.layout = Layout()


def _axes(chart: Axed, found: Plotted) -> None:
    horizontal = found.spec.kind == "barres_horizontales"
    categories, values = chart.x_axis, chart.y_axis
    # openpyxl 3.1 masque les axes par défaut (``delete``) : les montrer.
    categories.delete = False
    values.delete = False
    categories.number_format = "General"
    categories.majorTickMark = "none"
    categories.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=GRID, w=9525))
    categories.txPr = _text(900)
    values.number_format = AXIS_FORMATS[found.value_type]
    values.majorTickMark = "none"
    values.majorGridlines = ChartLines(
        spPr=GraphicalProperties(ln=LineProperties(solidFill=GRID, w=6350))
    )
    values.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
    values.txPr = _text(900)
    if horizontal:
        # La première catégorie en haut, comme dans le tableau ; l'axe des valeurs en bas.
        categories.scaling.orientation = "maxMin"
        values.crosses = "max"
    top = _top(found)
    if top is not None:
        values.scaling.min = 0
        values.scaling.max = top


def _top(found: Plotted) -> float | None:
    """Le bout de l'axe quand une étiquette doit tenir au bout des barres, comme dans Word.

    Seulement si toutes les valeurs sont connues : une formule n'a sa valeur
    qu'à l'ouverture, Excel choisit alors l'axe.
    """
    room = LABEL_ROOM.get(found.spec.kind)
    if room is None or not _labelled(found) or not found.complete:
        return None
    end = axis_end(found.known, room)
    return None if end is None else round(end, 10)
